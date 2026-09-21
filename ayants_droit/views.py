# ayants_droit/views.py
from io import BytesIO
import openpyxl
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import HttpResponse

from .models import AyantDroit
from .forms import AyantDroitForm
from core.mixins import enregistrer_audit

LIEN_PARENTE_MAP = {
    'CONJOINT': 'CONJOINT', 'CONJOINTE': 'CONJOINT', 'EPOUX': 'CONJOINT',
    'EPOUSE': 'CONJOINT', 'MARI': 'CONJOINT', 'FEMME': 'CONJOINT',
    'ENFANT': 'ENFANT', 'FILS': 'ENFANT', 'FILLE': 'ENFANT',
    'ENFANT ADOPTE': 'ENFANT_ADOPTE', 'ENFANT_ADOPTE': 'ENFANT_ADOPTE',
    'ADOPTE': 'ENFANT_ADOPTE',
    'PARENT': 'PARENT', 'PARENT A CHARGE': 'PARENT',
    'PERE': 'PARENT', 'MERE': 'PARENT',
    'FRERE': 'FRERE_SOEUR', 'SOEUR': 'FRERE_SOEUR',
    'FRERE_SOEUR': 'FRERE_SOEUR', 'FRERE/SOEUR': 'FRERE_SOEUR',
    'AUTRE': 'AUTRE',
}


@login_required
def telecharger_modele(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Modèle import"
    headers = [
        'N_affiliation', 'Nom', 'Prenom', 'Sexe', 'Date_naissance',
        'Lien_parente', 'Age_limite', 'Telephone', 'Statut',
        'Numero_identite', 'Notes',
    ]
    ws.append(headers)
    ws.append(['AFF-001', 'NDIMUBANDI', 'Marie', 'F', '1992-08-21',
               'CONJOINTE', '18', '79912345', 'ACTIF', '', ''])

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.font = openpyxl.styles.Font(bold=True, color='FFFFFF')
        cell.fill = openpyxl.styles.PatternFill(start_color='0a5c36', end_color='0a5c36', fill_type='solid')
        cell.alignment = openpyxl.styles.Alignment(horizontal='center', wrap_text=True)

    ws.freeze_panes = 'A2'
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="modele_import_ayants_droit.xlsx"'
    return response


@login_required
def importer(request):
    if request.user.est_super_admin:
        messages.warning(request, "⛔ Les Super Admins ne peuvent pas gérer les données d'une mutualité spécifique.")
        return redirect('superadmin:dashboard')
    if not request.user.peut_modifier:
        messages.error(request, "⛔ Permission refusée.")
        return redirect('ayants_droit:liste')

    from affilies.models import Affilie
    import datetime

    result = None

    if request.method == 'POST':
        affilie_id = request.POST.get('affilie')
        fichier = request.FILES.get('fichier')

        if not fichier:
            messages.error(request, "❌ Veuillez sélectionner un fichier Excel.")
        else:
            try:
                wb = openpyxl.load_workbook(fichier, read_only=True)
                ws = wb.active
                rows = list(ws.iter_rows(values_only=True))
                if len(rows) < 2:
                    raise ValueError("Le fichier doit contenir un en-tête et au moins une ligne de données.")

                headers = [str(h).strip().lower() if h else '' for h in rows[0]]
                mapping = {
                    'n_affiliation': 'affilie', 'numero affiliation': 'affilie',
                    'numero_affiliation': 'affilie', 'affiliation': 'affilie',
                    'id': 'affilie', 'id membre': 'affilie',
                    'nom': 'nom', 'prenom': 'prenom', 'sexe': 'sexe',
                    'date naissance': 'date_naissance', 'date_naissance': 'date_naissance',
                    'lien parente': 'lien_parente', 'lien_parente': 'lien_parente',
                    'parente': 'lien_parente',
                    'age limite': 'age_limite', 'age_limite': 'age_limite',
                    'telephone': 'telephone', 'statut': 'statut',
                    'numero identite': 'numero_identite', 'numero_identite': 'numero_identite',
                    'notes': 'notes',
                }

                col_map = {}
                for i, h in enumerate(headers):
                    if h in mapping:
                        col_map[mapping[h]] = i

                required = ['nom', 'sexe', 'date_naissance', 'lien_parente']
                missing_cols = [f for f in required if f not in col_map]
                if missing_cols:
                    raise ValueError(f"Colonnes obligatoires manquantes : {', '.join(missing_cols)}")

                affilie_unique = None
                if affilie_id:
                    affilie_unique = Affilie.objects.filter(
                        pk=affilie_id, mutualite=request.user.mutualite, supprime=False
                    ).first()

                affilies_cache = {}

                def trouver_affilie(numero):
                    if not numero:
                        return affilie_unique
                    if numero not in affilies_cache:
                        affilies_cache[numero] = Affilie.objects.filter(
                            mutualite=request.user.mutualite, supprime=False,
                            numero_affiliation=str(numero).strip(),
                        ).first()
                    return affilies_cache[numero]

                if not affilie_unique and 'affilie' not in col_map:
                    raise ValueError("Sélectionnez un affilié OU incluez la colonne N_affiliation dans le fichier.")

                succes, erreurs = 0, []

                for idx, row in enumerate(rows[1:], start=2):
                    try:
                        def val(field):
                            c = col_map.get(field)
                            return row[c] if c is not None and c < len(row) and row[c] is not None else None

                        affilie = trouver_affilie(val('affilie'))
                        if not affilie:
                            erreurs.append(f"Ligne {idx} : Affilié introuvable ({val('affilie') or 'aucun'})")
                            continue

                        nom = val('nom')
                        sexe = val('sexe')
                        date_naiss = val('date_naissance')
                        lien = val('lien_parente')

                        if not nom or not str(nom).strip():
                            erreurs.append(f"Ligne {idx} : Nom manquant")
                            continue
                        if not sexe or str(sexe).strip().upper() not in ('M', 'F'):
                            erreurs.append(f"Ligne {idx} : Sexe invalide (M ou F attendu)")
                            continue

                        lien_clean = str(lien).strip().upper()
                        lien_parente = LIEN_PARENTE_MAP.get(lien_clean)
                        if not lien_parente:
                            erreurs.append(f"Ligne {idx} : Lien de parenté invalide ({lien})")
                            continue

                        if date_naiss:
                            if isinstance(date_naiss, datetime.datetime):
                                date_naiss = date_naiss.date()
                            elif isinstance(date_naiss, str):
                                for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
                                    try:
                                        date_naiss = datetime.datetime.strptime(date_naiss.strip(), fmt).date()
                                        break
                                    except ValueError:
                                        pass
                                else:
                                    raise ValueError("Format date invalide")
                        else:
                            erreurs.append(f"Ligne {idx} : Date de naissance manquante")
                            continue

                        age_limite = val('age_limite')
                        try:
                            age_limite = int(age_limite) if age_limite not in (None, '') else 18
                        except (ValueError, TypeError):
                            age_limite = 18

                        statut_val = str(val('statut') or 'ACTIF').strip().upper()
                        if statut_val not in dict(AyantDroit.STATUT_CHOICES):
                            statut_val = 'ACTIF'

                        ad = AyantDroit(
                            mutualite=request.user.mutualite,
                            affilie=affilie,
                            nom=str(nom).strip().upper(),
                            prenom=str(val('prenom') or '').strip().capitalize(),
                            sexe=str(sexe).strip().upper(),
                            date_naissance=date_naiss,
                            lien_parente=lien_parente,
                            age_limite=age_limite,
                            telephone=str(val('telephone') or '').strip(),
                            statut=statut_val,
                            numero_identite=str(val('numero_identite') or '').strip(),
                            notes=str(val('notes') or '').strip(),
                            date_ajout=datetime.date.today(),
                        )
                        ad.save()
                        succes += 1
                    except Exception as e:
                        erreurs.append(f"Ligne {idx} : {str(e)}")

                result = {'succes': succes, 'erreurs': erreurs}
                if succes:
                    messages.success(request, f"✅ {succes} ayant(s) droit importé(s) avec succès.")
            except Exception as e:
                messages.error(request, f"❌ Erreur d'import : {str(e)}")

    from affilies.models import Affilie
    affilies = Affilie.objects.filter(mutualite=request.user.mutualite, supprime=False).order_by('numero_affiliation')
    return render(request, 'ayants_droit/importer.html', {
        'affilies': affilies,
        'result': result,
        'lien_choices': AyantDroit.LIEN_CHOICES,
    })


@login_required
def export_excel(request):
    if request.user.est_super_admin:
        messages.warning(request, "⛔ Les Super Admins ne peuvent pas gérer les données d'une mutualité spécifique.")
        return redirect('superadmin:dashboard')

    qs = AyantDroit.objects.filter(
        mutualite=request.user.mutualite, supprime=False,
    ).select_related('affilie')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ayants droit"
    headers = [
        'Affilié', 'N° Affiliation', 'Nom', 'Prénom', 'Sexe',
        'Date Naissance', 'Lien de parenté', 'Âge limite', 'Téléphone',
        'Statut', "Date d'ajout",
    ]
    ws.append(headers)
    for obj in qs:
        ws.append([
            obj.affilie.nom_complet if obj.affilie else '',
            obj.affilie.numero_affiliation if obj.affilie else '',
            obj.nom, obj.prenom, obj.get_sexe_display(),
            obj.date_naissance.strftime('%d/%m/%Y') if obj.date_naissance else '',
            obj.get_lien_parente_display(),
            obj.age_limite, obj.telephone, obj.get_statut_display(),
            obj.date_ajout.strftime('%d/%m/%Y') if obj.date_ajout else '',
        ])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="ayants_droit.xlsx"'
    return response


@login_required
def liste(request):
    if request.user.est_super_admin:
        messages.warning(request, "⛔ Les Super Admins ne peuvent pas gérer les données d'une mutualité spécifique.")
        return redirect('superadmin:dashboard')

    from .models import radier_majeurs
    radier_majeurs()

    qs = AyantDroit.objects.filter(mutualite=request.user.mutualite, supprime=False).select_related('affilie')

    q = request.GET.get('q', '')
    lien = request.GET.get('lien', '')
    statut = request.GET.get('statut', '')

    if q:
        qs = qs.filter(
            Q(nom__icontains=q) | Q(prenom__icontains=q) |
            Q(affilie__nom__icontains=q) | Q(affilie__numero_affiliation__icontains=q)
        )
    if lien:
        qs = qs.filter(lien_parente=lien)
    if statut:
        qs = qs.filter(statut=statut)

    paginator = Paginator(qs, 10)
    page = request.GET.get('page', 1)
    ayants = paginator.get_page(page)

    return render(request, 'ayants_droit/liste.html', {
        'ayants': ayants,
        'q': q,
        'lien': lien,
        'statut': statut,
        'total': qs.count(),
        'lien_choices': AyantDroit.LIEN_CHOICES,
        'statut_choices': AyantDroit.STATUT_CHOICES,
    })


@login_required
def ajouter(request):
    if request.user.est_super_admin:
        messages.warning(request, "⛔ Les Super Admins ne peuvent pas gérer les données d'une mutualité spécifique.")
        return redirect('superadmin:dashboard')
    if not request.user.peut_modifier:
        messages.error(request, "⛔ Vous êtes en mode consultation.")
        return redirect('ayants_droit:liste')

    # Pré-sélectionner l'affilié si passé en paramètre
    initial = {}
    affilie_id = request.GET.get('affilie')
    if affilie_id:
        initial['affilie'] = affilie_id

    form = AyantDroitForm(request.POST or None, user=request.user, initial=initial)
    if request.method == 'POST' and form.is_valid():
        obj = form.save(commit=False)
        obj.mutualite = request.user.mutualite
        obj.save()
        enregistrer_audit(request, 'CREATION', 'ayants_droit',
                          f"Création ayant droit : {obj.nom_complet}", 'AyantDroit', obj.pk)
        messages.success(request, f"✅ Ayant droit {obj.nom_complet} ajouté avec succès.")
        return redirect('ayants_droit:liste')

    return render(request, 'ayants_droit/form.html', {'form': form, 'titre': 'Ajouter un ayant droit'})


@login_required
def modifier(request, pk):
    if request.user.est_super_admin:
        messages.warning(request, "⛔ Les Super Admins ne peuvent pas gérer les données d'une mutualité spécifique.")
        return redirect('superadmin:dashboard')
    if not request.user.peut_modifier:
        messages.error(request, "⛔ Vous êtes en mode consultation.")
        return redirect('ayants_droit:liste')

    obj = get_object_or_404(AyantDroit, pk=pk, supprime=False, mutualite=request.user.mutualite)
    form = AyantDroitForm(request.POST or None, instance=obj, user=request.user)
    if request.method == 'POST' and form.is_valid():
        form.save()
        enregistrer_audit(request, 'MODIFICATION', 'ayants_droit',
                          f"Modification ayant droit : {obj.nom_complet}", 'AyantDroit', obj.pk)
        messages.success(request, f"✅ Ayant droit {obj.nom_complet} modifié.")
        return redirect('ayants_droit:liste')

    return render(request, 'ayants_droit/form.html', {'form': form, 'obj': obj, 'titre': 'Modifier ayant droit'})


@login_required
def supprimer(request, pk):
    if request.user.est_super_admin:
        messages.warning(request, "⛔ Les Super Admins ne peuvent pas gérer les données d'une mutualité spécifique.")
        return redirect('superadmin:dashboard')
    if not request.user.peut_supprimer:
        messages.error(request, "⛔ Permission insuffisante.")
        return redirect('ayants_droit:liste')

    obj = get_object_or_404(AyantDroit, pk=pk, is_deleted=False, mutualite=request.user.mutualite)
    if request.method == 'POST':
        nom = obj.nom_complet
        obj.delete(user=request.user)
        enregistrer_audit(request, 'SUPPRESSION', 'ayants_droit',
                          f"Suppression ayant droit : {nom}", 'AyantDroit', pk)
        messages.success(request, f"🗑️ {nom} déplacé vers la corbeille.")
        return redirect('ayants_droit:liste')

    return render(request, 'ayants_droit/confirm_supprimer.html', {'obj': obj})
