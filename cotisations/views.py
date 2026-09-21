# cotisations/views.py
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm, cm
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from io import BytesIO
from datetime import datetime

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Q, Sum
from django.utils import timezone
from django.http import HttpResponse

from .models import Cotisation
from .forms import CotisationForm
from core.mixins import enregistrer_audit
from affilies.models import Affilie

@login_required
def liste(request):
    if request.user.est_super_admin:
        messages.warning(request, "⛔ Les Super Admins ne peuvent pas gérer les données d'une mutualité spécifique.")
        return redirect('superadmin:dashboard')
    qs = Cotisation.objects.filter(mutualite=request.user.mutualite).select_related('affilie')
    
    q = request.GET.get('q', '')
    statut = request.GET.get('statut', '')
    annee = request.GET.get('annee', '')
    
    if q:
        qs = qs.filter(
            Q(reference__icontains=q) | 
            Q(affilie__nom__icontains=q) |
            Q(affilie__numero_affiliation__icontains=q)
        )
    if statut:
        qs = qs.filter(statut=statut)
    if annee:
        qs = qs.filter(annee=annee)
        
    paginator = Paginator(qs, 10)
    page = request.GET.get('page', 1)
    cotisations = paginator.get_page(page)
    
    return render(request, 'cotisations/liste.html', {
        'cotisations': cotisations,
        'q': q,
        'statut': statut,
        'annee': annee,
        'total': qs.count(),
        'statut_choices': Cotisation.STATUT_CHOICES,
    })

@login_required
def ajouter(request):
    if request.user.est_super_admin:
        messages.warning(request, "⛔ Les Super Admins ne peuvent pas gérer les données d'une mutualité spécifique.")
        return redirect('superadmin:dashboard')
    if not request.user.peut_modifier:
        messages.error(request, "⛔ Permission refusée.")
        return redirect('cotisations:liste')
        
    if request.method == 'POST':
        form = CotisationForm(request.POST, request.FILES, mutualite=request.user.mutualite)
        if form.is_valid():
            obj = form.save(commit=False)
            if not obj.reference or Cotisation._base_manager.filter(reference=obj.reference).exists():
                obj.reference = Cotisation.generer_reference()
            obj.enregistre_par = request.user
            obj.mutualite = request.user.mutualite
            obj.save()
            enregistrer_audit(request, 'CREATION', 'cotisations', f"Paiement cotisation : {obj.reference}", 'Cotisation', obj.pk)
            messages.success(request, f"✅ Cotisation {obj.reference} enregistrée avec succès.")
            return redirect('cotisations:liste')
    else:
        form = CotisationForm(initial={'reference': Cotisation.generer_reference()}, mutualite=request.user.mutualite)
        
    return render(request, 'cotisations/form.html', {'form': form, 'titre': 'Enregistrer une cotisation'})

@login_required
def modifier(request, pk):
    if request.user.est_super_admin:
        messages.warning(request, "⛔ Les Super Admins ne peuvent pas gérer les données d'une mutualité spécifique.")
        return redirect('superadmin:dashboard')
    if not request.user.peut_modifier:
        messages.error(request, "⛔ Permission refusée.")
        return redirect('cotisations:liste')
        
    obj = get_object_or_404(Cotisation, pk=pk, mutualite=request.user.mutualite)
    if request.method == 'POST':
        form = CotisationForm(request.POST, request.FILES, instance=obj, mutualite=request.user.mutualite)
        if form.is_valid():
            form.save()
            enregistrer_audit(request, 'MODIFICATION', 'cotisations', f"Modification cotisation : {obj.reference}", 'Cotisation', obj.pk)
            messages.success(request, f"✅ Cotisation {obj.reference} modifiée.")
            return redirect('cotisations:liste')
    else:
        form = CotisationForm(instance=obj, mutualite=request.user.mutualite)
        
    return render(request, 'cotisations/form.html', {'form': form, 'obj': obj, 'titre': 'Modifier la cotisation'})

@login_required
def supprimer(request, pk):
    if request.user.est_super_admin:
        messages.warning(request, "⛔ Les Super Admins ne peuvent pas gérer les données d'une mutualité spécifique.")
        return redirect('superadmin:dashboard')
    if not request.user.peut_supprimer:
        messages.error(request, "⛔ Permission refusée.")
        return redirect('cotisations:liste')
        
    obj = get_object_or_404(Cotisation, pk=pk, mutualite=request.user.mutualite)
    if request.method == 'POST':
        reference = obj.reference
        obj.delete(user=request.user)
        enregistrer_audit(request, 'SUPPRESSION', 'cotisations', f"Suppression cotisation : {reference}", 'Cotisation', pk)
        messages.success(request, f"🗑️ Cotisation {reference} supprimée.")
        return redirect('cotisations:liste')
        
    return render(request, 'cotisations/confirm_supprimer.html', {'obj': obj})


@login_required
def detail(request, pk):
    """Détail d'une cotisation."""
    if request.user.est_super_admin:
        messages.warning(request, "⛔ Les Super Admins ne peuvent pas gérer les données d'une mutualité spécifique.")
        return redirect('superadmin:dashboard')
    cotisation = get_object_or_404(Cotisation.objects.select_related('affilie', 'enregistre_par'), pk=pk, mutualite=request.user.mutualite)
    return render(request, 'cotisations/detail.html', {'cotisation': cotisation})


@login_required
def journal(request):
    """Journal mensuel des cotisations payées — matrice affiliés × mois."""
    if request.user.est_super_admin:
        messages.warning(request, "⛔ Les Super Admins ne peuvent pas gérer les données d'une mutualité spécifique.")
        return redirect('superadmin:dashboard')
    now = timezone.now()
    annee = int(request.GET.get('annee', now.year))
    q = request.GET.get('q', '').strip()

    qs = Cotisation.objects.filter(
        mutualite=request.user.mutualite,
        statut='PAYE',
        annee=annee,
        mois__isnull=False,
    ).select_related('affilie')

    rows = {}
    for c in qs:
        aff = c.affilie
        pk = aff.pk
        if pk not in rows:
            rows[pk] = {
                'affilie': aff,
                'mois': [0] * 13,
            }
        rows[pk]['mois'][c.mois] += c.montant

    lignes = sorted(rows.values(), key=lambda r: (r['affilie'].nom, r['affilie'].prenom))
    for r in lignes:
        r['total'] = sum(r['mois'][1:])

    if q:
        q_lower = q.lower()
        lignes = [
            r for r in lignes
            if q_lower in r['affilie'].numero_affiliation.lower()
            or (r['affilie'].numero_carte_identite and q_lower in r['affilie'].numero_carte_identite.lower())
            or q_lower in r['affilie'].nom.lower()
            or q_lower in r['affilie'].prenom.lower()
        ]

    total_montant = qs.aggregate(total=Sum('montant'))['total'] or 0
    nb_paiements = qs.count()

    total_par_mois = [0] * 13
    for r in lignes:
        for i in range(1, 13):
            total_par_mois[i] += r['mois'][i]

    paginator = Paginator(lignes, 15)
    page = request.GET.get('page', 1)
    lignes_page = paginator.get_page(page)

    annees = list(range(now.year, now.year - 5, -1))
    mois_noms = [
        'Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin',
        'Juillet', 'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre',
    ]

    return render(request, 'cotisations/journal.html', {
        'lignes': lignes_page,
        'total_montant': total_montant,
        'nb_paiements': nb_paiements,
        'annee': annee,
        'q': q,
        'total_par_mois': total_par_mois,
        'mois': '',
        'annees': annees,
        'mois_noms': mois_noms,
    })


@login_required
def valider_paiement(request, pk):
    """Marquer une cotisation comme payée."""
    if request.user.est_super_admin:
        messages.warning(request, "⛔ Les Super Admins ne peuvent pas gérer les données d'une mutualité spécifique.")
        return redirect('superadmin:dashboard')
    if not request.user.peut_valider:
        messages.error(request, "⛔ Permission refusée.")
        return redirect('cotisations:liste')

    obj = get_object_or_404(Cotisation, pk=pk, mutualite=request.user.mutualite)
    if request.method == 'POST':
        obj.statut = 'PAYE'
        obj.date_paiement = timezone.now().date()
        obj.save(update_fields=['statut', 'date_paiement'])
        enregistrer_audit(request, 'VALIDATION', 'cotisations', f"Validation paiement : {obj.reference}", 'Cotisation', pk)
        messages.success(request, f"✅ Cotisation {obj.reference} marquée comme payée.")
    return redirect('cotisations:liste')


@login_required
def export_excel(request):
    """Export des cotisations en Excel."""
    if request.user.est_super_admin:
        messages.warning(request, "⛔ Les Super Admins ne peuvent pas gérer les données d'une mutualité spécifique.")
        return redirect('superadmin:dashboard')
    if not request.user.peut_exporter:
        messages.error(request, "⛔ Permission refusée.")
        return redirect('cotisations:liste')

    qs = Cotisation.objects.filter(mutualite=request.user.mutualite).select_related('affilie')
    
    q = request.GET.get('q', '')
    statut = request.GET.get('statut', '')
    annee = request.GET.get('annee', '')
    
    if q:
        qs = qs.filter(
            Q(reference__icontains=q) | 
            Q(affilie__nom__icontains=q) |
            Q(affilie__numero_affiliation__icontains=q)
        )
    if statut:
        qs = qs.filter(statut=statut)
    if annee:
        qs = qs.filter(annee=annee)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cotisations"
    
    headers = [
        'Référence', 'N° Affiliation', 'Nom', 'Prénom', 
        'Période', 'Mois', 'Année', 'Montant (BIF)',
        'Statut', 'Mode de paiement', 'Date émission',
        'Date échéance', 'Date paiement', 'Date enregistrement'
    ]
    ws.append(headers)
    
    # Style en-tête
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='0a5c36', end_color='0a5c36', fill_type='solid')
    header_alignment = Alignment(horizontal='center', wrap_text=True)
    
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    for obj in qs:
        ws.append([
            obj.reference,
            obj.affilie.numero_affiliation,
            obj.affilie.nom,
            obj.affilie.prenom,
            obj.get_periode_type_display(),
            obj.mois or '',
            obj.annee,
            float(obj.montant),
            obj.get_statut_display(),
            obj.get_mode_paiement_display() if obj.mode_paiement else '',
            obj.date_emission.strftime('%d/%m/%Y') if obj.date_emission else '',
            obj.date_echeance.strftime('%d/%m/%Y') if obj.date_echeance else '',
            obj.date_paiement.strftime('%d/%m/%Y') if obj.date_paiement else '',
            obj.date_creation.strftime('%d/%m/%Y %H:%M') if obj.date_creation else '',
        ])
    
    # Ajuster la largeur des colonnes
    column_widths = [18, 18, 20, 18, 15, 8, 8, 15, 15, 15, 15, 15, 15, 20]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"cotisations_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@login_required
def export_pdf(request):
    """Export des cotisations en PDF."""
    if request.user.est_super_admin:
        messages.warning(request, "⛔ Les Super Admins ne peuvent pas gérer les données d'une mutualité spécifique.")
        return redirect('superadmin:dashboard')
    if not request.user.peut_exporter:
        messages.error(request, "⛔ Permission refusée.")
        return redirect('cotisations:liste')

    qs = Cotisation.objects.filter(mutualite=request.user.mutualite).select_related('affilie')
    
    q = request.GET.get('q', '')
    statut = request.GET.get('statut', '')
    annee = request.GET.get('annee', '')
    
    if q:
        qs = qs.filter(
            Q(reference__icontains=q) | 
            Q(affilie__nom__icontains=q) |
            Q(affilie__numero_affiliation__icontains=q)
        )
    if statut:
        qs = qs.filter(statut=statut)
    if annee:
        qs = qs.filter(annee=annee)

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=1*cm,
        rightMargin=1*cm,
        topMargin=1.5*cm,
        bottomMargin=1.5*cm,
    )
    
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='TitleFR', fontName='Helvetica-Bold', fontSize=14, alignment=TA_CENTER, spaceAfter=6, textColor=colors.HexColor('#0a5c36')))
    styles.add(ParagraphStyle(name='SubTitleFR', fontName='Helvetica', fontSize=9, alignment=TA_CENTER, spaceAfter=12, textColor=colors.grey))
    styles.add(ParagraphStyle(name='CellStyle', fontName='Helvetica', fontSize=7, leading=9))
    styles.add(ParagraphStyle(name='HeaderStyle', fontName='Helvetica-Bold', fontSize=7, leading=9, textColor=colors.white, alignment=TA_CENTER))
    styles.add(ParagraphStyle(name='TotalStyle', fontName='Helvetica-Bold', fontSize=8, leading=10, textColor=colors.HexColor('#0a5c36'), alignment=TA_RIGHT))
    
    elements = []
    
    # En-tête
    mutualite = request.user.mutualite
    elements.append(Paragraph(f"<b>{mutualite.nom}</b>", styles['TitleFR']))
    elements.append(Paragraph(f"Liste des cotisations - {timezone.now().strftime('%d/%m/%Y à %H:%M')}", styles['SubTitleFR']))
    elements.append(Spacer(1, 10))
    
    # Données du tableau
    headers = [
        'Référence', 'N° Affil.', 'Nom', 'Prénom', 
        'Période', 'Mois', 'Année', 'Montant',
        'Statut', 'Paiement', 'Émission', 'Échéance', 'Paiement'
    ]
    
    data = [headers]
    total_montant = 0
    
    for obj in qs:
        total_montant += float(obj.montant)
        data.append([
            obj.reference,
            obj.affilie.numero_affiliation,
            obj.affilie.nom,
            obj.affilie.prenom,
            obj.get_periode_type_display(),
            str(obj.mois) if obj.mois else '',
            str(obj.annee),
            f"{float(obj.montant):,.0f}",
            obj.get_statut_display(),
            obj.get_mode_paiement_display() if obj.mode_paiement else '',
            obj.date_emission.strftime('%d/%m/%Y') if obj.date_emission else '',
            obj.date_echeance.strftime('%d/%m/%Y') if obj.date_echeance else '',
            obj.date_paiement.strftime('%d/%m/%Y') if obj.date_paiement else '',
        ])
    
    # Ligne de total
    data.append(['', '', '', '', '', '', '', f"{total_montant:,.0f}", 'TOTAL', '', '', '', ''])
    
    # Créer le tableau
    col_widths = [2.2*cm, 1.8*cm, 2*cm, 1.8*cm, 1.5*cm, 0.8*cm, 0.8*cm, 2*cm, 1.5*cm, 1.5*cm, 1.5*cm, 1.5*cm, 1.5*cm]
    
    table = Table(data, colWidths=col_widths, repeatRows=1)
    
    style_cmds = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0a5c36')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 7),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('FONTNAME', (0, 1), (-1, -2), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -2), 7),
        ('ALIGN', (7, 1), (7, -1), 'RIGHT'),
        ('ALIGN', (5, 1), (6, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -2), 0.5, colors.lightgrey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#f5f9f5')]),
        # Ligne de total
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#e8f5e9')),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, -1), (-1, -1), 8),
        ('TEXTCOLOR', (0, -1), (-1, -1), colors.HexColor('#0a5c36')),
        ('LINEABOVE', (0, -1), (-1, -1), 1.5, colors.HexColor('#0a5c36')),
    ]
    
    table.setStyle(TableStyle(style_cmds))
    elements.append(table)
    
    # Pied de page
    elements.append(Spacer(1, 10))
    elements.append(Paragraph(f"Total : {qs.count()} cotisation(s) - Montant total : {total_montant:,.0f} BIF", styles['TotalStyle']))
    elements.append(Spacer(1, 6))
    elements.append(Paragraph(f"Document généré par {request.user.nom_complet or request.user.username} le {timezone.now().strftime('%d/%m/%Y à %H:%M')}", styles['SubTitleFR']))
    
    doc.build(elements)
    buffer.seek(0)
    
    filename = f"cotisations_{timezone.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    response = HttpResponse(buffer.read(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@login_required
def telecharger_modele_import(request):
    """Télécharger le modèle Excel pour l'import de cotisations."""
    if request.user.est_super_admin:
        messages.warning(request, "⛔ Les Super Admins ne peuvent pas gérer les données d'une mutualité spécifique.")
        return redirect('superadmin:dashboard')
    if not request.user.peut_modifier:
        messages.error(request, "⛔ Permission refusée.")
        return redirect('cotisations:liste')
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Modèle import cotisations"
    
    headers = [
        'Nom', 'Prénom', 'Montant', 'Date_paiement', 'Annee', 'Periode_type',
        'Mois', 'Mode_paiement', 'Reference', 'Date_emission', 'Date_echeance', 'Notes'
    ]
    ws.append(headers)
    
    # Ligne d'exemple
    ws.append([
        'NDIMUBANDI', 'Jean', '15000', '2026-07-20', '2026', 'MENSUEL',
        '7', 'CASH', '', '2026-07-01', '2026-07-31', 'Cotisation juillet 2026'
    ])
    
    # Style en-tête
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='0a5c36', end_color='0a5c36', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
    
    # Largeurs de colonnes
    column_widths = [20, 18, 12, 15, 8, 15, 8, 15, 18, 15, 15, 30]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="modele_import_cotisations.xlsx"'
    return response


@login_required
def importer_excel(request):
    """Importer des cotisations depuis un fichier Excel."""
    if request.user.est_super_admin:
        messages.warning(request, "⛔ Les Super Admins ne peuvent pas gérer les données d'une mutualité spécifique.")
        return redirect('superadmin:dashboard')
    if not request.user.peut_modifier:
        messages.error(request, "⛔ Permission refusée.")
        return redirect('cotisations:liste')
    
    result = None
    
    if request.method == 'POST':
        fichier = request.FILES.get('fichier')
        # Récupérer le type de période depuis le formulaire (fallback)
        periode_type_form = request.POST.get('periode_type_import', '').strip().upper()
        
        if not fichier:
            messages.error(request, "❌ Veuillez sélectionner un fichier Excel.")
        else:
            try:
                wb = openpyxl.load_workbook(fichier, read_only=True)
                ws = wb.active
                rows = list(ws.iter_rows(values_only=True))
                
                if len(rows) < 2:
                    raise ValueError("Le fichier doit contenir un en-tête et au moins une ligne de données.")
                
                headers = [str(h).strip().lower() if h else '' for h in rows[0]]
                mapping = {
                    'nom': 'nom', 'prenom': 'prenom', 'montant': 'montant',
                    'date_paiement': 'date_paiement', 'annee': 'annee', 'periode_type': 'periode_type',
                    'mois': 'mois', 'mode_paiement': 'mode_paiement', 'reference': 'reference',
                    'date_emission': 'date_emission', 'date_echeance': 'date_echeance', 'notes': 'notes',
                }
                
                col_map = {}
                for i, h in enumerate(headers):
                    if h in mapping:
                        col_map[mapping[h]] = i
                
                required = ['nom', 'prenom', 'montant', 'date_paiement']
                missing_cols = [f for f in required if f not in col_map]
                if missing_cols:
                    raise ValueError(f"Colonnes obligatoires manquantes : {', '.join(missing_cols)}")
                
                succes, erreurs = 0, []
                
                for idx, row in enumerate(rows[1:], start=2):
                    try:
                        def val(field):
                            c = col_map.get(field)
                            return row[c] if c is not None and c < len(row) and row[c] is not None else None
                        
                        nom = val('nom')
                        prenom = val('prenom')
                        montant = val('montant')
                        date_paiement = val('date_paiement')
                        annee = val('annee')
                        periode_type = val('periode_type')
                        mois = val('mois')
                        mode_paiement = val('mode_paiement')
                        reference = val('reference')
                        date_emission = val('date_emission')
                        date_echeance = val('date_echeance')
                        notes = val('notes')
                        
                        if not nom or not str(nom).strip():
                            erreurs.append(f"Ligne {idx} : Nom manquant")
                            continue
                        if not prenom or not str(prenom).strip():
                            erreurs.append(f"Ligne {idx} : Prénom manquant")
                            continue
                        
                        # Trouver l'affilié
                        nom_clean = str(nom).strip().upper()
                        prenom_clean = str(prenom).strip().capitalize()
                        affilie = Affilie.objects.filter(
                            mutualite=request.user.mutualite,
                            nom__iexact=nom_clean,
                            prenom__iexact=prenom_clean
                        ).first()
                        
                        if not affilie:
                            erreurs.append(f"Ligne {idx} : Affilié '{nom_clean} {prenom_clean}' introuvable")
                            continue
                        
                        # Valider le montant
                        try:
                            montant = float(montant) if montant else 0
                        except (ValueError, TypeError):
                            erreurs.append(f"Ligne {idx} : Montant invalide")
                            continue
                        
                        if montant <= 0:
                            erreurs.append(f"Ligne {idx} : Montant doit être > 0")
                            continue
                        
                        # Valider date_paiement
                        if date_paiement:
                            if isinstance(date_paiement, datetime):
                                date_paiement = date_paiement.date()
                            elif isinstance(date_paiement, str):
                                parsed = False
                                for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
                                    try:
                                        date_paiement = datetime.strptime(date_paiement.strip(), fmt).date()
                                        parsed = True
                                        break
                                    except ValueError:
                                        pass
                                if not parsed:
                                    erreurs.append(f"Ligne {idx} : Format date_paiement invalide (attendu YYYY-MM-DD ou DD/MM/YYYY)")
                                    continue
                        else:
                            erreurs.append(f"Ligne {idx} : Date de paiement manquante")
                            continue
                        
                        # Année
                        try:
                            annee = int(annee) if annee else date_paiement.year
                        except (ValueError, TypeError):
                            annee = date_paiement.year
                        
                        # Période type
                        periode_choices = dict(Cotisation.PERIODE_CHOICES)
                        if periode_type and str(periode_type).strip().upper() in periode_choices:
                            periode_type = str(periode_type).strip().upper()
                        elif periode_type_form and periode_type_form in periode_choices:
                            periode_type = periode_type_form
                        else:
                            periode_type = 'MENSUEL'
                        
                        # Mois
                        try:
                            mois = int(mois) if mois else date_paiement.month
                            if not (1 <= mois <= 12):
                                mois = date_paiement.month
                        except (ValueError, TypeError):
                            mois = date_paiement.month
                        
                        # Mode paiement
                        mode_choices = dict(Cotisation.MODE_PAIEMENT_CHOICES)
                        if mode_paiement and str(mode_paiement).strip().upper() in mode_choices:
                            mode_paiement = str(mode_paiement).strip().upper()
                        else:
                            mode_paiement = 'CASH'
                        
                        # Date émission
                        if date_emission:
                            if isinstance(date_emission, datetime):
                                date_emission = date_emission.date()
                            elif isinstance(date_emission, str):
                                for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
                                    try:
                                        date_emission = datetime.strptime(date_emission.strip(), fmt).date()
                                        break
                                    except ValueError:
                                        pass
                        else:
                            date_emission = date_paiement
                        
                        # Date échéance
                        if date_echeance:
                            if isinstance(date_echeance, datetime):
                                date_echeance = date_echeance.date()
                            elif isinstance(date_echeance, str):
                                for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
                                    try:
                                        date_echeance = datetime.strptime(date_echeance.strip(), fmt).date()
                                        break
                                    except ValueError:
                                        pass
                        
                        # Créer la cotisation
                        cot = Cotisation(
                            mutualite=request.user.mutualite,
                            affilie=affilie,
                            reference=reference or Cotisation.generer_reference(),
                            montant=montant,
                            periode_type=periode_type,
                            mois=mois,
                            annee=annee,
                            date_emission=date_emission,
                            date_echeance=date_echeance,
                            date_paiement=date_paiement,
                            statut='PAYE',
                            mode_paiement=mode_paiement,
                            notes=notes or '',
                            enregistre_par=request.user,
                        )
                        cot.save()
                        enregistrer_audit(request, 'CREATION', 'cotisations', f"Import cotisation : {cot.reference}", 'Cotisation', cot.pk)
                        succes += 1
                        
                    except Exception as e:
                        erreurs.append(f"Ligne {idx} : {str(e)}")
                
                result = {'succes': succes, 'erreurs': erreurs}
                if succes:
                    messages.success(request, f"✅ {succes} cotisation(s) importée(s) avec succès.")
                if erreurs:
                    for err in erreurs[:10]:
                        messages.error(request, err)
                    if len(erreurs) > 10:
                        messages.error(request, f"... et {len(erreurs) - 10} autres erreurs.")
                        
            except Exception as e:
                messages.error(request, f"❌ Erreur d'import : {str(e)}")
    
    return render(request, 'cotisations/importer.html', {'result': result})


# ── Rapport Périodique (Cotisations & Factures par Affilié / Mois) ──

MOIS_NOMS_FR = {
    1: 'Jan', 2: 'Fév', 3: 'Mar', 4: 'Avr',
    5: 'Mai', 6: 'Juin', 7: 'Juil', 8: 'Août',
    9: 'Sep', 10: 'Oct', 11: 'Nov', 12: 'Déc',
}

MOIS_NOMS_LONGS = {
    1: 'Janvier', 2: 'Février', 3: 'Mars', 4: 'Avril',
    5: 'Mai', 6: 'Juin', 7: 'Juillet', 8: 'Août',
    9: 'Septembre', 10: 'Octobre', 11: 'Novembre', 12: 'Décembre',
}

def _construire_donnees_rapport_periodique(mutualite, md, ad, mf, af, origine_id=None, statut_filter=None, q=None):
    from datetime import date
    from calendar import monthrange
    from factures.models import Facture
    
    # 1. Générer la liste des mois
    date_debut = date(ad, md, 1)
    dernier_jour = monthrange(af, mf)[1]
    date_fin = date(af, mf, dernier_jour)
    
    mois_list = []
    curr = date_debut
    while curr <= date_fin:
        mois_list.append((curr.month, curr.year, f"{MOIS_NOMS_FR[curr.month]} {curr.year}"))
        if curr.month == 12:
            curr = date(curr.year + 1, 1, 1)
        else:
            curr = date(curr.year, curr.month + 1, 1)
        if len(mois_list) >= 36:  # Sécurité max 36 mois
            break
            
    # 2. Récupérer les affiliés
    qs_aff = Affilie.objects.filter(mutualite=mutualite, supprime=False).select_related('origine', 'categorie_origine')
    if origine_id:
        qs_aff = qs_aff.filter(origine_id=origine_id)
    if statut_filter == 'ACTIF':
        qs_aff = qs_aff.filter(statut='ACTIF')
    if q:
        qs_aff = qs_aff.filter(
            Q(nom__icontains=q) |
            Q(prenom__icontains=q) |
            Q(numero_affiliation__icontains=q)
        )
    affilies = list(qs_aff.order_by('nom', 'prenom'))
    affilie_ids = [a.id for a in affilies]

    # 3. Récupérer et indexer les cotisations (statut PAYE)
    cot_dict = {}  # (affilie_id, mois, annee) -> montant
    if affilie_ids:
        cots = Cotisation.objects.filter(
            mutualite=mutualite,
            affilie_id__in=affilie_ids,
            statut='PAYE'
        )
        for c in cots:
            m = c.mois or (c.date_paiement.month if c.date_paiement else c.date_emission.month)
            a = c.annee or (c.date_paiement.year if c.date_paiement else c.date_emission.year)
            key = (c.affilie_id, m, a)
            cot_dict[key] = cot_dict.get(key, 0) + float(c.montant)

    # 4. Récupérer et indexer les factures (statuts VALIDE ou PAYE)
    fac_dict = {}  # (affilie_id, mois, annee) -> montant_mutualite
    if affilie_ids:
        factures = Facture.objects.filter(
            mutualite=mutualite,
            statut__in=['VALIDE', 'PAYE'],
            date__gte=date_debut,
            date__lte=date_fin
        ).select_related('ayant_droit')
        
        for f in factures:
            aff_id = f.affilie_id or (f.ayant_droit.affilie_id if f.ayant_droit else None)
            if aff_id and aff_id in affilie_ids:
                key = (aff_id, f.date.month, f.date.year)
                montant = float(f.montant_mutualite if f.montant_mutualite is not None else f.montant_total)
                fac_dict[key] = fac_dict.get(key, 0) + montant

    # 5. Construire la matrice de données
    rows = []
    totaux_mois = [{'mois': m, 'annee': a, 'label': lbl, 'total_cotisation': 0.0, 'total_facture': 0.0, 'solde': 0.0} for m, a, lbl in mois_list]
    grand_total_cotisations = 0.0
    grand_total_factures = 0.0

    for aff in affilies:
        mois_data = []
        tot_cot = 0.0
        tot_fac = 0.0
        
        for idx, (m, a, lbl) in enumerate(mois_list):
            cot_val = cot_dict.get((aff.id, m, a), 0.0)
            fac_val = fac_dict.get((aff.id, m, a), 0.0)
            diff = cot_val - fac_val
            
            mois_data.append({
                'mois': m,
                'annee': a,
                'cotisation': cot_val,
                'facture': fac_val,
                'solde': diff,
            })
            
            tot_cot += cot_val
            tot_fac += fac_val
            totaux_mois[idx]['total_cotisation'] += cot_val
            totaux_mois[idx]['total_facture'] += fac_val
            totaux_mois[idx]['solde'] += diff

        solde_aff = tot_cot - tot_fac
        ratio_conso = round((tot_fac / tot_cot * 100), 1) if tot_cot > 0 else (100.0 if tot_fac > 0 else 0.0)
        
        rows.append({
            'affilie': aff,
            'mois_data': mois_data,
            'total_cotisations': tot_cot,
            'total_factures': tot_fac,
            'solde': solde_aff,
            'ratio_conso': ratio_conso,
        })
        
        grand_total_cotisations += tot_cot
        grand_total_factures += tot_fac

    grand_total_solde = grand_total_cotisations - grand_total_factures
    taux_sinistralite = round((grand_total_factures / grand_total_cotisations * 100), 1) if grand_total_cotisations > 0 else 0.0

    return {
        'mois_list': mois_list,
        'rows': rows,
        'totaux_mois': totaux_mois,
        'grand_total_cotisations': grand_total_cotisations,
        'grand_total_factures': grand_total_factures,
        'grand_total_solde': grand_total_solde,
        'taux_sinistralite': taux_sinistralite,
        'nb_affilies': len(affilies),
        'date_debut': date_debut,
        'date_fin': date_fin,
    }


@login_required
def rapport_periodique(request):
    if request.user.est_super_admin:
        messages.warning(request, "⛔ Les Super Admins ne peuvent pas gérer les données d'une mutualité spécifique.")
        return redirect('superadmin:dashboard')
    
    from .forms import RapportPeriodiqueFilterForm
    now = timezone.now()
    
    # Valeurs par défaut : 6 derniers mois de l'année en cours
    init_md = max(1, now.month - 5)
    init_ad = now.year
    init_mf = now.month
    init_af = now.year

    md = int(request.GET.get('mois_debut', init_md))
    ad = int(request.GET.get('annee_debut', init_ad))
    mf = int(request.GET.get('mois_fin', init_mf))
    af = int(request.GET.get('annee_fin', init_af))
    origine_id = request.GET.get('origine')
    statut_filter = request.GET.get('statut', 'ACTIF')
    q = request.GET.get('q', '').strip()

    form = RapportPeriodiqueFilterForm(
        initial={
            'mois_debut': md,
            'annee_debut': ad,
            'mois_fin': mf,
            'annee_fin': af,
            'origine': origine_id,
            'statut': statut_filter,
            'q': q,
        },
        mutualite=request.user.mutualite
    )

    data = _construire_donnees_rapport_periodique(
        mutualite=request.user.mutualite,
        md=md, ad=ad, mf=mf, af=af,
        origine_id=origine_id,
        statut_filter=statut_filter,
        q=q
    )

    # Pagination 20 par page
    paginator = Paginator(data['rows'], 20)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)

    return render(request, 'cotisations/rapport_periodique.html', {
        'form': form,
        'data': data,
        'page_obj': page_obj,
        'md': md, 'ad': ad, 'mf': mf, 'af': af,
        'origine_id': origine_id,
        'statut_filter': statut_filter,
        'q': q,
    })


@login_required
def rapport_periodique_excel(request):
    if request.user.est_super_admin:
        return redirect('superadmin:dashboard')
        
    now = timezone.now()
    md = int(request.GET.get('mois_debut', 1))
    ad = int(request.GET.get('annee_debut', now.year))
    mf = int(request.GET.get('mois_fin', now.month))
    af = int(request.GET.get('annee_fin', now.year))
    origine_id = request.GET.get('origine')
    statut_filter = request.GET.get('statut', 'ACTIF')
    q = request.GET.get('q', '').strip()

    data = _construire_donnees_rapport_periodique(
        mutualite=request.user.mutualite,
        md=md, ad=ad, mf=mf, af=af,
        origine_id=origine_id,
        statut_filter=statut_filter,
        q=q
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rapport Périodique"

    # Styles
    font_title = Font(name='Segoe UI', size=14, bold=True, color='1B3B2B')
    font_header = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
    font_bold = Font(name='Segoe UI', size=9, bold=True)
    font_regular = Font(name='Segoe UI', size=9)
    fill_header = PatternFill(start_color='1B3B2B', end_color='1B3B2B', fill_type='solid')
    fill_sub_cot = PatternFill(start_color='D1E7DD', end_color='D1E7DD', fill_type='solid')
    fill_sub_fac = PatternFill(start_color='F8D7DA', end_color='F8D7DA', fill_type='solid')
    fill_total = PatternFill(start_color='F3F4F6', end_color='F3F4F6', fill_type='solid')
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')

    # Titre
    mut_nom = request.user.mutualite.nom if request.user.mutualite else "Mutualité de Santé"
    ws.merge_cells('A1:G1')
    ws['A1'] = f"{mut_nom} — Rapport Périodique des Cotisations & Soins"
    ws['A1'].font = font_title
    ws['A1'].alignment = align_left

    ws.merge_cells('A2:G2')
    ws['A2'] = f"Période du {data['date_debut'].strftime('%d/%m/%Y')} au {data['date_fin'].strftime('%d/%m/%Y')} — Généré le {now.strftime('%d/%m/%Y à %H:%M')}"
    ws['A2'].font = Font(name='Segoe UI', size=9, italic=True, color='6B7280')

    # Ligne 4 : Entêtes groupées
    row_num = 4
    ws.cell(row=row_num, column=1, value="N° Affiliation").font = font_header
    ws.cell(row=row_num, column=1).fill = fill_header
    ws.cell(row=row_num, column=1).alignment = align_center

    ws.cell(row=row_num, column=2, value="Nom & Prénom").font = font_header
    ws.cell(row=row_num, column=2).fill = fill_header
    ws.cell(row=row_num, column=2).alignment = align_center

    ws.cell(row=row_num, column=3, value="Structure / Origine").font = font_header
    ws.cell(row=row_num, column=3).fill = fill_header
    ws.cell(row=row_num, column=3).alignment = align_center

    col_idx = 4
    for m, a, lbl in data['mois_list']:
        ws.merge_cells(start_row=row_num, start_column=col_idx, end_row=row_num, end_column=col_idx + 1)
        cell = ws.cell(row=row_num, column=col_idx, value=lbl)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        col_idx += 2

    # Totaux globaux colonnes
    ws.cell(row=row_num, column=col_idx, value="Total Cotisations").font = font_header
    ws.cell(row=row_num, column=col_idx).fill = fill_header
    ws.cell(row=row_num, column=col_idx).alignment = align_center
    col_idx += 1

    ws.cell(row=row_num, column=col_idx, value="Total Soins").font = font_header
    ws.cell(row=row_num, column=col_idx).fill = fill_header
    ws.cell(row=row_num, column=col_idx).alignment = align_center
    col_idx += 1

    ws.cell(row=row_num, column=col_idx, value="Solde (BIF)").font = font_header
    ws.cell(row=row_num, column=col_idx).fill = fill_header
    ws.cell(row=row_num, column=col_idx).alignment = align_center

    # Ligne 5 : Sous-entêtes (Cotis. / Soins)
    row_num = 5
    ws.cell(row=row_num, column=1, value="")
    ws.cell(row=row_num, column=2, value="")
    ws.cell(row=row_num, column=3, value="")
    sub_col = 4
    for _ in data['mois_list']:
        c1 = ws.cell(row=row_num, column=sub_col, value="Cotis.")
        c1.font = font_bold
        c1.fill = fill_sub_cot
        c1.alignment = align_center

        c2 = ws.cell(row=row_num, column=sub_col + 1, value="Soins")
        c2.font = font_bold
        c2.fill = fill_sub_fac
        c2.alignment = align_center
        sub_col += 2

    ws.cell(row=row_num, column=sub_col, value="").fill = fill_total
    ws.cell(row=row_num, column=sub_col + 1, value="").fill = fill_total
    ws.cell(row=row_num, column=sub_col + 2, value="").fill = fill_total

    # Données lignes affiliés
    row_num = 6
    for r in data['rows']:
        aff = r['affilie']
        ws.cell(row=row_num, column=1, value=aff.numero_affiliation).font = font_regular
        ws.cell(row=row_num, column=2, value=aff.nom_complet).font = font_bold
        ws.cell(row=row_num, column=3, value=aff.origine.nom if aff.origine else "-").font = font_regular

        cur_col = 4
        for m_data in r['mois_data']:
            c_cell = ws.cell(row=row_num, column=cur_col, value=m_data['cotisation'])
            c_cell.font = font_regular
            c_cell.alignment = align_right
            c_cell.number_format = '#,##0'

            f_cell = ws.cell(row=row_num, column=cur_col + 1, value=m_data['facture'])
            f_cell.font = font_regular
            f_cell.alignment = align_right
            f_cell.number_format = '#,##0'
            cur_col += 2

        tot_c = ws.cell(row=row_num, column=cur_col, value=r['total_cotisations'])
        tot_c.font = font_bold
        tot_c.alignment = align_right
        tot_c.number_format = '#,##0'

        tot_f = ws.cell(row=row_num, column=cur_col + 1, value=r['total_factures'])
        tot_f.font = font_bold
        tot_f.alignment = align_right
        tot_f.number_format = '#,##0'

        solde_c = ws.cell(row=row_num, column=cur_col + 2, value=r['solde'])
        solde_c.font = font_bold
        solde_c.alignment = align_right
        solde_c.number_format = '#,##0'
        if r['solde'] < 0:
            solde_c.font = Font(name='Segoe UI', size=9, bold=True, color='DC2626')

        row_num += 1

    # Ligne TOTAUX GLOBAUX
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=3)
    tot_label = ws.cell(row=row_num, column=1, value="TOTAUX GÉNÉRAUX")
    tot_label.font = font_bold
    tot_label.alignment = Alignment(horizontal='center', vertical='center')
    tot_label.fill = fill_total

    cur_col = 4
    for tm in data['totaux_mois']:
        tc = ws.cell(row=row_num, column=cur_col, value=tm['total_cotisation'])
        tc.font = font_bold
        tc.alignment = align_right
        tc.number_format = '#,##0'
        tc.fill = fill_total

        tf = ws.cell(row=row_num, column=cur_col + 1, value=tm['total_facture'])
        tf.font = font_bold
        tf.alignment = align_right
        tf.number_format = '#,##0'
        tf.fill = fill_total
        cur_col += 2

    gtc = ws.cell(row=row_num, column=cur_col, value=data['grand_total_cotisations'])
    gtc.font = font_bold
    gtc.alignment = align_right
    gtc.number_format = '#,##0'
    gtc.fill = fill_total

    gtf = ws.cell(row=row_num, column=cur_col + 1, value=data['grand_total_factures'])
    gtf.font = font_bold
    gtf.alignment = align_right
    gtf.number_format = '#,##0'
    gtf.fill = fill_total

    gts = ws.cell(row=row_num, column=cur_col + 2, value=data['grand_total_solde'])
    gts.font = font_bold
    gts.alignment = align_right
    gts.number_format = '#,##0'
    gts.fill = fill_total

    # Ajustement largeur des colonnes
    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 26
    ws.column_dimensions['C'].width = 22
    for c_i in range(4, cur_col + 3):
        col_letter = openpyxl.utils.get_column_letter(c_i)
        ws.column_dimensions[col_letter].width = 14

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="Rapport_Periodique_{data["date_debut"].strftime("%Y%m")}_{data["date_fin"].strftime("%Y%m")}.xlsx"'
    wb.save(response)
    return response

