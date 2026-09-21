import logging
from datetime import datetime as timezone
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.shortcuts import get_object_or_404, render, redirect
from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_POST
from django.views.generic import DetailView
from django.views.generic.edit import FormMixin
from core.views import CrudListView, CrudCreateView, CrudUpdateView, CrudDeleteView, enregistrer_audit

from .models import Affilie
from .forms import AffilieForm
from django.contrib import messages
from django.db.models import Q
import openpyxl
from io import BytesIO

logger = logging.getLogger('core.security')

FORMATS_NUMERO = {
    '0-1': 1,
    '00-01': 2,
    '000-001': 3,
    '0000-0001': 4,
    '00000-00001': 5,
}
PREFIXE_AFFILIATION = 'AFF-'


def prochain_numero_affiliation():
    """Retourne le prochain numéro séquentiel global (toutes mutualités) pour le préfixe AFF-."""
    dernier = 0
    for num in Affilie._base_manager.exclude(supprime=True).filter(
            numero_affiliation__startswith=PREFIXE_AFFILIATION
    ).values_list('numero_affiliation', flat=True):
        try:
            dernier = max(dernier, int(num[len(PREFIXE_AFFILIATION):]))
        except ValueError:
            pass
    return dernier + 1


def formater_numero_affiliation(numero, largeur):
    """Formate un numéro séquentiel avec padding de zéros et le préfixe AFF-."""
    return f"{PREFIXE_AFFILIATION}{numero:0{largeur}d}"


def largeur_affiliation_defaut():
    """Largeur du format en vigueur (dernier AFF- créé), ou 1 (0-1) si aucun import n'existe."""
    for num in Affilie._base_manager.exclude(supprime=True).filter(
            numero_affiliation__startswith=PREFIXE_AFFILIATION
    ).order_by('-id').values_list('numero_affiliation', flat=True)[:100]:
        suffixe = num[len(PREFIXE_AFFILIATION):]
        if suffixe.isdigit():
            return max(1, min(len(suffixe), 5))
    return 1


def format_affiliation_defaut():
    """Clé du format par défaut ('0-1', '00-01', ...)."""
    largeur = largeur_affiliation_defaut()
    for fmt, w in FORMATS_NUMERO.items():
        if w == largeur:
            return fmt
    return '0-1'


class AffilieListView(CrudListView):
    model = Affilie
    context_object_name = 'affilies'
    module_name = 'affilies'
    search_fields = ['nom', 'prenom', 'numero_affiliation', 'telephone']
    filter_fields = {'statut': 'statut', 'origine': 'origine_id'}
    default_order = '-date_creation'
    list_template = 'affilies/liste.html'
    titre = 'Liste des affiliés'

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx['statut_choices'] = Affilie.STATUT_CHOICES
        from origines.models import Origine
        ctx['origines'] = Origine.objects.filter(mutualite=self.request.user.mutualite, actif=True)
        return ctx


class AffilieCreateView(CrudCreateView):
    model = Affilie
    form_class = AffilieForm
    module_name = 'affilies'
    form_template = 'affilies/form.html'
    titre = 'Ajouter un affilié'
    success_message_create = "✅ Affilié {obj.nom_complet} ajouté avec succès."

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['mutualite'] = self.request.user.mutualite
        return kwargs

    def get_initial(self):
        from django.utils import timezone
        initial = super().get_initial()
        initial['date_adhesion'] = timezone.localdate()
        initial['numero_affiliation'] = formater_numero_affiliation(
            prochain_numero_affiliation(), largeur_affiliation_defaut()
        )
        return initial

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        from origines.models import Origine, CategorieOrigine
        import json
        plafond_fields = [
            'cotisation_mensuelle', 'taux_prise_en_charge', 'soins_ambulatoires',
            'limite_laboratoire', 'limite_kinesitherapie', 'limite_reanimation', 'limite_imagerie_medicale',
            'hopital_plafond_chambre', 'hpc_nuit', 'limite_hospitalisation_generale',
            'maternite_forfait', 'pourcentage_maternite',
            'pharmacie_forfait', 'limites_medicaments', 'limites_lunettes', 'limites_dents',
        ]
        # Origines
        origines = Origine.objects.filter(mutualite=self.request.user.mutualite)
        data_origines = {}
        for o in origines:
            data_origines[str(o.pk)] = {f: str(getattr(o, f, '0')) for f in plafond_fields}
        ctx['origin_plafonds_json'] = json.dumps(data_origines)
        
        # Catégories
        categories = CategorieOrigine.objects.filter(mutualite=self.request.user.mutualite)
        data_categories = {}
        for c in categories:
            data_categories[str(c.pk)] = {f: str(getattr(c, f, '0')) for f in plafond_fields}
        ctx['categorie_plafonds_json'] = json.dumps(data_categories)
        return ctx

    def before_save(self, obj):
        obj.cree_par = self.request.user
        # Sécurité : s'assurer que la mutualité est définie
        if not obj.mutualite_id:
            obj.mutualite = self.request.user.mutualite
        plafond_fields = [
            'cotisation_mensuelle', 'taux_prise_en_charge', 'soins_ambulatoires',
            'limite_laboratoire', 'limite_kinesitherapie', 'limite_reanimation', 'limite_imagerie_medicale',
            'hopital_plafond_chambre', 'hpc_nuit', 'limite_hospitalisation_generale',
            'maternite_forfait', 'pourcentage_maternite',
            'pharmacie_forfait', 'limites_medicaments', 'limites_lunettes', 'limites_dents',
        ]
        # Priorité 1 : plafonds de la catégorie d'origine (si définie)
        if obj.categorie_origine_id:
            for field in plafond_fields:
                val = getattr(obj.categorie_origine, field, None)
                if val:
                    setattr(obj, field, val)
        # Priorité 2 : plafonds de l'origine (si définie)
        elif obj.origine_id:
            for field in plafond_fields:
                val = getattr(obj.origine, field, None)
                if val:
                    setattr(obj, field, val)
        # Priorité 3 : plafonds par défaut de la mutualité
        elif obj.mutualite:
            m = obj.mutualite
            plafonds_map = {
                'cotisation_mensuelle': m.def_cotisation_mensuelle,
                'soins_ambulatoires': m.def_soins_ambulatoires,
                'limite_hospitalisation_generale': m.def_limite_hospitalisation_generale,
                'hopital_plafond_chambre': m.def_hopital_plafond_chambre,
                'hpc_nuit': m.def_hpc_nuit,
                'limite_laboratoire': m.def_limite_laboratoire,
                'limite_imagerie_medicale': m.def_limite_imagerie_medicale,
                'limite_kinesitherapie': m.def_limite_kinesitherapie,
                'limite_reanimation': m.def_limite_reanimation,
                'maternite_forfait': m.def_maternite_forfait,
                'pourcentage_maternite': m.def_pourcentage_maternite,
                'pharmacie_forfait': m.def_pharmacie_forfait,
                'limites_medicaments': m.def_limites_medicaments,
                'limites_lunettes': m.def_limites_lunettes,
                'limites_dents': m.def_limites_dents,
            }
            for field, value in plafonds_map.items():
                if value:
                    setattr(obj, field, value)


class AffilieDetailView(LoginRequiredMixin, DetailView):
    model = Affilie
    template_name = 'affilies/detail.html'
    context_object_name = 'affilie'

    def get_object(self, queryset=None):
        return get_object_or_404(Affilie, pk=self.kwargs['pk'], supprime=False)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        affilie = self.object
        ctx['ayants_droit'] = []
        ctx['cotisations'] = []
        try:
            from ayants_droit.models import radier_majeurs
            radier_majeurs()
            ctx['ayants_droit'] = affilie.ayants_droit.filter(supprime=False)
        except Exception as e:
            logger.warning("Erreur chargement ayants_droit: %s", e)
        try:
            ctx['cotisations'] = affilie.cotisations.all().order_by('-date_paiement')[:10]
        except Exception as e:
            logger.warning("Erreur chargement cotisations: %s", e)

        # Exclusions
        try:
            from affilies.models import AffilieExclusion
            from origines.models import OrigineExclusion
            from core.models import MutualiteExclusion
            medicaments = affilie.mutualite.medicaments.all() if affilie.mutualite else []
            services = affilie.mutualite.servicesantes.all() if affilie.mutualite else []
            ctx['medicaments'] = medicaments
            ctx['services'] = services
            ctx['exclusions'] = AffilieExclusion.objects.filter(affilie=affilie).select_related('medicament', 'service')
            ctx['exclusions_globales'] = MutualiteExclusion.objects.filter(mutualite=affilie.mutualite).select_related('medicament', 'service') if affilie.mutualite else []
            if affilie.origine:
                ctx['exclusions_origine'] = OrigineExclusion.objects.filter(origine=affilie.origine).select_related('medicament', 'service')
            else:
                ctx['exclusions_origine'] = []
        except Exception as e:
            logger.warning("Erreur chargement exclusions: %s", e)
            ctx['exclusions'] = []
            ctx['exclusions_globales'] = []
            ctx['exclusions_origine'] = []

        return ctx


class AffilieUpdateView(CrudUpdateView):
    model = Affilie
    form_class = AffilieForm
    module_name = 'affilies'
    form_template = 'affilies/form.html'
    titre = "Modifier l'affilié"
    success_message_update = "✅ Affilié {obj.nom_complet} modifié avec succès."
    success_url = None

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['mutualite'] = self.request.user.mutualite
        return kwargs


class AffilieDeleteView(CrudDeleteView):
    model = Affilie
    module_name = 'affilies'
    confirm_delete_template = 'affilies/confirm_supprimer.html'
    success_message_delete = "🗑️ {obj} a été déplacé vers la corbeille."

    def get_success_url(self):
        from django.urls import reverse
        return reverse('affilies:liste')


@login_required
def telecharger_modele(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Modèle import"
    headers = [
        'Nom', 'Prenom', 'Sexe', 'Date_naissance',
        'Etat_civil', 'Numero_carte_identite',
        'Province', 'Commune', 'Zone', 'Colline', 'Adresse',
        'Telephone', 'Email', 'Profession',
        'Groupe_sanguin', 'Contact_urgence_nom', 'Contact_urgence_telephone',
        'Centre_sante_reference', 'Hopital_reference',
        'Date_adhesion', 'Type_adhesion', 'Section_antenne',
        'Categorie', 'Statut', 'Notes',
    ]
    ws.append(headers)
    ws.append(['NDIMUBANDI', 'Jean', 'M', '1990-05-15', 'MARIE',
               '123456', 'BUJUMBURA', 'MUKAZA', 'CENTRE', 'Bwiza', 'Av. Large',
               '79123456', 'jean@mail.com', 'Enseignant',
               'B+', 'MARIE NDIMUBANDI', '79987654',
               'CS Kinama', 'HPRC',
               '2026-07-20', 'INDIVIDUELLE', 'Section Nord',
               'Cadre', 'ACTIF', 'Exemple'])

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.font = openpyxl.styles.Font(bold=True, color='FFFFFF')
        cell.fill = openpyxl.styles.PatternFill(start_color='0a5c36', end_color='0a5c36', fill_type='solid')
        cell.alignment = openpyxl.styles.Alignment(horizontal='center', wrap_text=True)

    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 8
    ws.column_dimensions['E'].width = 15

    from django.http import HttpResponse
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="modele_import_affilies.xlsx"'
    return response


@login_required
def importer(request):
    if request.user.est_super_admin:
        messages.warning(request, "⛔ Les Super Admins ne peuvent pas gérer les données d'une mutualité spécifique.")
        return redirect('superadmin:dashboard')
    if not request.user.peut_modifier:
        messages.error(request, "⛔ Permission refusée.")
        return redirect('affilies:liste')

    from origines.models import Origine
    import datetime

    result = None
    format_defaut = format_affiliation_defaut()

    if request.method == 'POST':
        origine_id = request.POST.get('origine')
        fichier = request.FILES.get('fichier')

        if not fichier:
            messages.error(request, "❌ Veuillez sélectionner un fichier Excel.")
        elif not origine_id:
            messages.error(request, "❌ Veuillez sélectionner une origine.")
        else:
            try:
                wb = openpyxl.load_workbook(fichier, read_only=True)
                ws = wb.active
                rows = list(ws.iter_rows(values_only=True))
                if len(rows) < 2:
                    raise ValueError("Le fichier doit contenir un en-tête et au moins une ligne de données.")

                headers = [str(h).strip().lower() if h else '' for h in rows[0]]
                mapping = {
                    'nom': 'nom', 'prenom': 'prenom',
                    'sexe': 'sexe', 'date naissance': 'date_naissance', 'date_naissance': 'date_naissance',
                    'etat civil': 'etat_civil', 'etat_civil': 'etat_civil',
                    'numero carte identite': 'numero_carte_identite', 'numero_carte_identite': 'numero_carte_identite',
                    'province': 'province', 'commune': 'commune', 'zone': 'zone',
                    'colline': 'colline', 'adresse': 'adresse',
                    'telephone': 'telephone', 'email': 'email', 'profession': 'profession',
                    'groupe sanguin': 'groupe_sanguin', 'groupe_sanguin': 'groupe_sanguin',
                    'contact urgence nom': 'contact_urgence_nom', 'contact_urgence_nom': 'contact_urgence_nom',
                    'contact urgence telephone': 'contact_urgence_telephone', 'contact_urgence_telephone': 'contact_urgence_telephone',
                    'centre sante reference': 'centre_sante_reference', 'centre_sante_reference': 'centre_sante_reference',
                    'hopital reference': 'hopital_reference', 'hopital_reference': 'hopital_reference',
                    'date adhesion': 'date_adhesion', 'date_adhesion': 'date_adhesion',
                    'type adhesion': 'type_adhesion', 'type_adhesion': 'type_adhesion',
                    'section antenne': 'section_antenne', 'section_antenne': 'section_antenne',
                    'categorie': 'categorie', 'categorie employe': 'categorie', 'category': 'categorie',
                    'statut': 'statut', 'notes': 'notes',
                }

                col_map = {}
                for i, h in enumerate(headers):
                    if h in mapping:
                        col_map[mapping[h]] = i

                required = ['nom', 'prenom', 'sexe', 'date_naissance']
                missing_cols = [f for f in required if f not in col_map]
                if missing_cols:
                    raise ValueError(f"Colonnes obligatoires manquantes : {', '.join(missing_cols)}")

                origine = Origine.objects.get(pk=origine_id, mutualite=request.user.mutualite)
                from origines.models import CategorieOrigine as _Cat
                categorie_forcee = None
                _cat_id = (request.POST.get('categorie') or '').strip()
                if _cat_id:
                    categorie_forcee = _Cat.objects.filter(pk=_cat_id, origine=origine, mutualite=request.user.mutualite).first()
                    if not categorie_forcee:
                        raise ValueError("Catégorie invalide pour cette origine.")
                succes, erreurs = 0, []

                format_numero = request.POST.get('format_numero') or format_defaut
                largeur = FORMATS_NUMERO.get(format_numero, 1)
                prochain_numero = prochain_numero_affiliation()

                for idx, row in enumerate(rows[1:], start=2):
                    try:
                        def val(field):
                            c = col_map.get(field)
                            return row[c] if c is not None and c < len(row) and row[c] is not None else None

                        nom = val('nom')
                        prenom = val('prenom')
                        sexe = val('sexe')
                        date_naiss = val('date_naissance')

                        if not nom or not str(nom).strip():
                            erreurs.append(f"Ligne {idx} : Nom manquant")
                            continue
                        if not prenom or not str(prenom).strip():
                            erreurs.append(f"Ligne {idx} : Prénom manquant")
                            continue
                        if not sexe or str(sexe).strip().upper() not in ('M', 'F'):
                            erreurs.append(f"Ligne {idx} : Sexe invalide (M ou F attendu)")
                            continue

                        if date_naiss:
                            if isinstance(date_naiss, datetime.datetime):
                                date_naiss = date_naiss.date()
                            elif isinstance(date_naiss, str):
                                for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
                                    try:
                                        date_naiss = datetime.datetime.strptime(date_naiss.strip(), fmt).date()
                                        break
                                    except ValueError:
                                        pass
                                else:
                                    raise ValueError("Format date invalide")
                        else:
                            erreurs.append(f"Ligne {idx} : Date de naissance manquante")
                            continue

                        aff = Affilie(mutualite=request.user.mutualite, origine=origine)
                        aff.nom = str(nom).strip().upper()
                        aff.prenom = str(prenom).strip().capitalize()
                        aff.sexe = str(sexe).strip().upper()
                        aff.date_naissance = date_naiss

                        optional_fields = [
                            ('numero_carte_identite', str, lambda v: str(v).strip()),
                            ('etat_civil', str, lambda v: str(v).strip().upper() if str(v).strip().upper() in dict(Affilie.ETAT_CIVIL_CHOICES) else 'CELIBATAIRE'),
                            ('province', str, lambda v: str(v).strip()),
                            ('commune', str, lambda v: str(v).strip()),
                            ('zone', str, lambda v: str(v).strip()),
                            ('colline', str, lambda v: str(v).strip()),
                            ('adresse', str, lambda v: str(v).strip()),
                            ('telephone', str, lambda v: str(v).strip()),
                            ('email', str, lambda v: str(v).strip()),
                            ('profession', str, lambda v: str(v).strip()),
                            ('groupe_sanguin', str, lambda v: str(v).strip().upper() if str(v).strip().upper() in dict(Affilie.GROUPE_SANGUIN_CHOICES) else ''),
                            ('contact_urgence_nom', str, lambda v: str(v).strip()),
                            ('contact_urgence_telephone', str, lambda v: str(v).strip()),
                            ('centre_sante_reference', str, lambda v: str(v).strip()),
                            ('hopital_reference', str, lambda v: str(v).strip()),
                            ('section_antenne', str, lambda v: str(v).strip()),
                            ('notes', str, lambda v: str(v).strip()),
                            ('statut', str, lambda v: str(v).strip().upper() if str(v).strip().upper() in dict(Affilie.STATUT_CHOICES) else 'ACTIF'),
                            ('type_adhesion', str, lambda v: str(v).strip().upper() if str(v).strip().upper() in dict(Affilie.TYPE_ADHESION_CHOICES) else 'INDIVIDUELLE'),
                        ]

                        for field, _, transform in optional_fields:
                            v = val(field)
                            if v:
                                try:
                                    setattr(aff, field, transform(v))
                                except:
                                    pass

                        date_adh = val('date_adhesion')
                        if date_adh:
                            if isinstance(date_adh, datetime.datetime):
                                aff.date_adhesion = date_adh.date()
                            elif isinstance(date_adh, str):
                                for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
                                    try:
                                        aff.date_adhesion = datetime.datetime.strptime(date_adh.strip(), fmt).date()
                                        break
                                    except ValueError:
                                        pass

                        # ── Catégorie : forcée par le formulaire OU colonne fichier ──
                        from origines.models import CategorieOrigine
                        _plafonds = ['cotisation_mensuelle', 'soins_ambulatoires',
                                      'limite_laboratoire', 'limite_kinesitherapie', 'limite_reanimation',
                                      'limite_imagerie_medicale', 'hopital_plafond_chambre', 'hpc_nuit',
                                      'limite_hospitalisation_generale', 'maternite_forfait', 'pourcentage_maternite',
                                      'pharmacie_forfait', 'limites_medicaments', 'limites_lunettes', 'limites_dents']
                        if categorie_forcee:
                            aff.categorie_origine = categorie_forcee
                            for f in _plafonds:
                                setattr(aff, f, getattr(categorie_forcee, f, 0))
                        else:
                            cat_nom = val('categorie')
                            if cat_nom and str(cat_nom).strip():
                                cat_nom = str(cat_nom).strip()
                                cat = CategorieOrigine.objects.filter(
                                    origine=origine, mutualite=request.user.mutualite,
                                    nom__iexact=cat_nom, actif=True,
                                ).first()
                                if not cat:
                                    cat = CategorieOrigine.objects.create(
                                        mutualite=request.user.mutualite, origine=origine,
                                        nom=cat_nom,
                                        cotisation_mensuelle=origine.cotisation_mensuelle,
                                        soins_ambulatoires=origine.soins_ambulatoires,
                                        limite_laboratoire=origine.limite_laboratoire,
                                        limite_kinesitherapie=origine.limite_kinesitherapie,
                                        limite_reanimation=origine.limite_reanimation,
                                        limite_imagerie_medicale=origine.limite_imagerie_medicale,
                                        hopital_plafond_chambre=origine.hopital_plafond_chambre,
                                        hpc_nuit=origine.hpc_nuit,
                                        limite_hospitalisation_generale=origine.limite_hospitalisation_generale,
                                        maternite_forfait=origine.maternite_forfait,
                                        pourcentage_maternite=origine.pourcentage_maternite,
                                        pharmacie_forfait=origine.pharmacie_forfait,
                                        limites_medicaments=origine.limites_medicaments,
                                        limites_lunettes=origine.limites_lunettes,
                                        limites_dents=origine.limites_dents,
                                    )
                                aff.categorie_origine = cat
                                for f in _plafonds:
                                    setattr(aff, f, getattr(cat, f, 0))
                            else:
                                aff.cotisation_mensuelle = 0

                        aff.numero_affiliation = formater_numero_affiliation(prochain_numero, largeur)
                        prochain_numero += 1
                        aff.save()
                        succes += 1
                    except Exception as e:
                        erreurs.append(f"Ligne {idx} : {str(e)}")

                result = {'succes': succes, 'erreurs': erreurs}
                if succes:
                    messages.success(request, f"✅ {succes} affilié(s) importé(s) avec succès.")
            except Exception as e:
                messages.error(request, f"❌ Erreur d'import : {str(e)}")

    origines = Origine.objects.filter(mutualite=request.user.mutualite, actif=True)
    import json
    from origines.models import CategorieOrigine
    cats_map = {}
    for c in CategorieOrigine.objects.filter(mutualite=request.user.mutualite, actif=True).select_related('origine').order_by('nom'):
        cats_map.setdefault(str(c.origine_id), []).append({'pk': c.pk, 'nom': c.nom, 'cotisation': str(c.cotisation_mensuelle)})
    return render(request, 'affilies/importer.html', {
        'origines': origines,
        'result': result,
        'format_defaut': format_defaut,
        'categories_map_json': json.dumps(cats_map),
    })


@login_required
def export_excel(request):
    if request.user.est_super_admin:
        messages.warning(request, "⛔ Les Super Admins ne peuvent pas gérer les données d'une mutualité spécifique.")
        return redirect('superadmin:dashboard')
    qs = Affilie.objects.filter(supprime=False, mutualite=request.user.mutualite)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Affiliés"
    headers = ['N° Affiliation', 'Nom', 'Prénom', 'Sexe', 'Date Naissance', 'Téléphone', 'Origine', 'Catégorie', 'Type adhésion', 'Statut']
    ws.append(headers)
    for obj in qs:
        ws.append([
            obj.numero_affiliation, obj.nom, obj.prenom, obj.get_sexe_display(),
            obj.date_naissance.strftime('%d/%m/%Y') if obj.date_naissance else '', obj.telephone,
            obj.origine.nom if obj.origine else '', obj.categorie_origine.nom if obj.categorie_origine else '',
            obj.get_type_adhesion_display(), obj.get_statut_display()
        ])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="affilies.xlsx"'
    return response


def fiche_pdf(request, pk):
    """Génère la fiche affilié en PDF."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm, cm
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    import os

    affilie = get_object_or_404(Affilie, pk=pk, supprime=False, mutualite=request.user.mutualite)

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=2*cm,
        rightMargin=2*cm,
        topMargin=2*cm,
        bottomMargin=2*cm,
    )

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='TitleFR', fontName='Helvetica-Bold', fontSize=16, alignment=TA_CENTER, spaceAfter=6, textColor=colors.HexColor('#0a5c36')))
    styles.add(ParagraphStyle(name='SubTitleFR', fontName='Helvetica', fontSize=10, alignment=TA_CENTER, spaceAfter=12, textColor=colors.grey))
    styles.add(ParagraphStyle(name='SectionTitle', fontName='Helvetica-Bold', fontSize=11, spaceBefore=12, spaceAfter=6, textColor=colors.HexColor('#0a5c36')))
    styles.add(ParagraphStyle(name='BodyFR', fontName='Helvetica', fontSize=9, leading=12))
    styles.add(ParagraphStyle(name='SmallFR', fontName='Helvetica', fontSize=8, leading=10, textColor=colors.grey))

    elements = []

    # En-tête avec logo
    header_data = []
    logo_path = None
    if affilie.mutualite.logo and affilie.mutualite.logo.path:
        if os.path.exists(affilie.mutualite.logo.path):
            logo_path = affilie.mutualite.logo.path
    if not logo_path:
        from core.models import ConfigurationApp
        config = ConfigurationApp.get_config()
        if config and config.logo_app and os.path.exists(config.logo_app.path):
            logo_path = config.logo_app.path

    if logo_path:
        try:
            img = Image(logo_path, width=3*cm, height=3*cm)
            header_data.append([img, Paragraph(f"<b>{affilie.mutualite.nom}</b><br/>Fiche Affilié", styles['TitleFR']), ''])
        except:
            header_data.append(['', Paragraph(f"<b>{affilie.mutualite.nom}</b><br/>Fiche Affilié", styles['TitleFR']), ''])
    else:
        header_data.append(['', Paragraph(f"<b>{affilie.mutualite.nom}</b><br/>Fiche Affilié", styles['TitleFR']), ''])

    header_table = Table(header_data, colWidths=[4*cm, 11*cm, 4*cm])
    header_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('ALIGN', (0,0), (0,0), 'LEFT'),
        ('ALIGN', (1,0), (1,0), 'CENTER'),
        ('ALIGN', (2,0), (2,0), 'RIGHT'),
    ]))
    elements.append(header_table)
    elements.append(Spacer(1, 10))

    # Photo + Infos principales
    photo_cell = ''
    if affilie.photo and affilie.photo.path and os.path.exists(affilie.photo.path):
        try:
            photo_cell = Image(affilie.photo.path, width=3*cm, height=3*cm)
        except:
            photo_cell = 'Photo non disponible'
    else:
        photo_cell = 'Pas de photo'

    info_principale = [
        ['Nom complet', affilie.nom_complet],
        ['N° Affiliation', affilie.numero_affiliation],
        ['Date de naissance', affilie.date_naissance.strftime('%d/%m/%Y') if affilie.date_naissance else '-'],
        ['Âge', f"{affilie.age} ans" if affilie.age else '-'],
        ['Sexe', affilie.get_sexe_display()],
        ['État civil', affilie.get_etat_civil_display() if affilie.etat_civil else '-'],
        ['Groupe sanguin', affilie.get_groupe_sanguin_display() if affilie.groupe_sanguin else '-'],
        ['Téléphone', affilie.telephone or '-'],
        ['Email', affilie.email or '-'],
        ['Profession', affilie.profession or '-'],
        ['Catégorie', affilie.get_type_adhesion_display()],
        ['Origine', affilie.origine.nom if affilie.origine else 'Individuelle'],
        ['Statut', affilie.get_statut_display()],
        ['Date adhésion', affilie.date_adhesion.strftime('%d/%m/%Y') if affilie.date_adhesion else '-'],
    ]

    data = [[photo_cell, Paragraph('<b>Informations principales</b>', styles['SectionTitle']), '']]
    for label, value in info_principale:
        data.append([Paragraph(label, styles['SmallFR']), Paragraph(str(value), styles['BodyFR']), ''])

    t = Table(data, colWidths=[3.5*cm, 10*cm, 4.5*cm])
    t.setStyle(TableStyle([
        ('SPAN', (0,0), (0,1)),
        ('VALIGN', (0,0), (0,0), 'TOP'),
        ('ALIGN', (0,0), (0,0), 'CENTER'),
        ('FONTNAME', (0,1), (-1,-1), 'Helvetica'),
        ('FONTSIZE', (0,1), (-1,-1), 9),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('GRID', (0,1), (1,-1), 0.5, colors.lightgrey),
        ('BACKGROUND', (0,1), (0,-1), colors.HexColor('#f5f9f5')),
    ]))
    elements.append(t)
    elements.append(Spacer(1, 12))

    # Localisation
    loc_data = [['Localisation', '', '']]
    for label, value in [
        ('Province', affilie.province or '-'),
        ('Commune', affilie.commune or '-'),
        ('Zone', affilie.zone or '-'),
        ('Colline/Quartier', affilie.colline or '-'),
        ('Adresse', affilie.adresse or '-'),
    ]:
        loc_data.append([Paragraph(label, styles['SmallFR']), Paragraph(value, styles['BodyFR']), ''])
    t2 = Table(loc_data, colWidths=[4*cm, 12*cm, 2*cm])
    t2.setStyle(TableStyle([
        ('SPAN', (0,0), (2,0)),
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#e8f5e9')),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 9),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('GRID', (0,0), (-1,-1), 0.5, colors.lightgrey),
        ('BACKGROUND', (0,1), (0,-1), colors.HexColor('#f5f9f5')),
    ]))
    elements.append(t2)
    elements.append(Spacer(1, 12))

    # Médical
    med_data = [['Informations médicales', '', '']]
    for label, value in [
        ('Contact urgence', f"{affilie.contact_urgence_nom or '-'} {f'({affilie.contact_urgence_telephone})' if affilie.contact_urgence_telephone else ''}"),
        ('Centre de santé réf.', affilie.centre_sante_reference or '-'),
        ('Hôpital de réf.', affilie.hopital_reference or '-'),
    ]:
        med_data.append([Paragraph(label, styles['SmallFR']), Paragraph(value, styles['BodyFR']), ''])
    t3 = Table(med_data, colWidths=[4*cm, 12*cm, 2*cm])
    t3.setStyle(TableStyle([
        ('SPAN', (0,0), (2,0)),
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#e8f5e9')),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 9),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('GRID', (0,0), (-1,-1), 0.5, colors.lightgrey),
        ('BACKGROUND', (0,1), (0,-1), colors.HexColor('#f5f9f5')),
    ]))
    elements.append(t3)
    elements.append(Spacer(1, 12))

    # Ayants droit
    ayants = affilie.ayants_droit.filter(supprime=False)
    if ayants.exists():
        ad_data = [['Ayants droit', '', '', '']]
        ad_data.append([Paragraph('Nom complet', styles['SmallFR']), Paragraph('Lien', styles['SmallFR']), Paragraph('Âge', styles['SmallFR']), Paragraph('Statut', styles['SmallFR'])])
        for ad in ayants:
            ad_data.append([
                Paragraph(ad.nom_complet, styles['BodyFR']),
                Paragraph(ad.get_lien_parente_display(), styles['BodyFR']),
                Paragraph(f"{ad.age} ans" if ad.age else '-', styles['BodyFR']),
                Paragraph(ad.get_statut_display(), styles['BodyFR']),
            ])
        t4 = Table(ad_data, colWidths=[6*cm, 4*cm, 2*cm, 4*cm])
        t4.setStyle(TableStyle([
            ('SPAN', (0,0), (3,0)),
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#e8f5e9')),
            ('BACKGROUND', (0,1), (-1,1), colors.HexColor('#d4edda')),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTNAME', (0,1), (-1,1), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,-1), 9),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('GRID', (0,0), (-1,-1), 0.5, colors.lightgrey),
        ]))
        elements.append(t4)
        elements.append(Spacer(1, 12))

    # Garanties / Plafonds
    garanties = [
        ('Cotisation mensuelle', 'cotisation_mensuelle'),
        ('Soins ambulatoires', 'soins_ambulatoires'),
        ('Laboratoire', 'limite_laboratoire'),
        ('Kinésithérapie', 'limite_kinesitherapie'),
        ('Réanimation', 'limite_reanimation'),
        ('Imagerie médicale', 'limite_imagerie_medicale'),
        ('Hôpital plafond chambre', 'hopital_plafond_chambre'),
        ('HPC / Nuit', 'hpc_nuit'),
        ('Hospitalisation générale', 'limite_hospitalisation_generale'),
        ('Maternité forfait', 'maternite_forfait'),
        ('% Maternité', 'pourcentage_maternite'),
        ('Pharmacie forfait', 'pharmacie_forfait'),
        ('Médicaments', 'limites_medicaments'),
        ('Lunettes', 'limites_lunettes'),
        ('Soins dentaires', 'limites_dents'),
    ]

    g_data = [['Garanties / Plafonds', 'Montant']]
    for label, field in garanties:
        value = getattr(affilie, field, None)
        if field == 'pourcentage_maternite':
            display = f"{value} %" if value is not None else '-'
        else:
            display = f"{value:,.0f} BIF" if value is not None else '-'
        g_data.append([Paragraph(label, styles['SmallFR']), Paragraph(display, styles['BodyFR'])])

    t5 = Table(g_data, colWidths=[10*cm, 6*cm])
    t5.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0a5c36')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 9),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('GRID', (0,0), (-1,-1), 0.5, colors.lightgrey),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f5f9f5')]),
        ('ALIGN', (1,0), (1,-1), 'RIGHT'),
    ]))
    elements.append(t5)
    elements.append(Spacer(1, 20))

    # Pied de page
    elements.append(Paragraph(f"Document généré le {timezone.now().strftime('%d/%m/%Y à %H:%M')} par {request.user.nom_complet or request.user.username}", styles['SmallFR']))

    doc.build(elements)
    buffer.seek(0)

    response = HttpResponse(buffer.read(), content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="fiche_{affilie.numero_affiliation}.pdf"'
    return response


@login_required
def export_pdf(request):
    return HttpResponse("Export PDF à implémenter")


@login_required
def affilie_details_ajax(request, pk):
    """AJAX endpoint pour récupérer les détails d'un affilié."""
    try:
        affilie = get_object_or_404(Affilie, pk=pk, mutualite=request.user.mutualite)
        data = {
            'nom': affilie.nom,
            'prenom': affilie.prenom,
            'nom_complet': affilie.nom_complet,
            'sexe': affilie.sexe,
            'date_naissance': affilie.date_naissance.strftime('%Y-%m-%d') if affilie.date_naissance else '',
            'telephone': affilie.telephone or '',
            'adresse': affilie.adresse or '',
            'numero_carte_identite': affilie.numero_carte_identite or '',
            'groupe_sanguin': affilie.groupe_sanguin or '',
            'contact_urgence_nom': affilie.contact_urgence_nom or '',
            'contact_urgence_telephone': affilie.contact_urgence_telephone or '',
            'centre_sante_reference': affilie.centre_sante_reference or '',
            'hopital_reference': affilie.hopital_reference or '',
        }
        return JsonResponse({'success': True, 'data': data})
    except Affilie.DoesNotExist:
        return JsonResponse({'success': False, 'error': 'Affilié non trouvé'}, status=404)
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
def origin_plafonds_ajax(request, pk):
    from origines.models import Origine
    try:
        origine = get_object_or_404(Origine, pk=pk, mutualite=request.user.mutualite)
        plafond_fields = [
            'cotisation_mensuelle', 'soins_ambulatoires',
            'limite_laboratoire', 'limite_kinesitherapie', 'limite_reanimation', 'limite_imagerie_medicale',
            'hopital_plafond_chambre', 'hpc_nuit', 'limite_hospitalisation_generale',
            'maternite_forfait', 'pourcentage_maternite',
            'pharmacie_forfait', 'limites_medicaments', 'limites_lunettes', 'limites_dents',
        ]
        data = {f: str(getattr(origine, f, '0')) for f in plafond_fields}
        return JsonResponse({'success': True, 'data': data})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


@login_required
def categories_par_origine_ajax(request, pk):
    from origines.models import CategorieOrigine
    cats = CategorieOrigine.objects.filter(
        origine_id=pk, mutualite=request.user.mutualite, actif=True,
    ).values('pk', 'nom', 'cotisation_mensuelle').order_by('nom')
    return JsonResponse({'success': True, 'data': list(cats)})


# ─────────────────────────────────────────────
#  Exclusions médicaments / services (affilié)
# ─────────────────────────────────────────────

@login_required
@require_POST
def exclusion_ajouter(request, affilie_pk):
    from affilies.models import AffilieExclusion
    affilie = get_object_or_404(Affilie, pk=affilie_pk, mutualite=request.user.mutualite)
    type_excl = request.POST.get('type')
    item_pk = request.POST.get('item_pk')
    motif = request.POST.get('motif', '')

    if not type_excl or not item_pk:
        messages.error(request, "Veuillez sélectionner un type et un élément.")
        return redirect('affilies:detail', pk=affilie_pk)

    if type_excl == 'medicament':
        from pharmacie.models import Medicament
        item = get_object_or_404(Medicament, pk=item_pk, mutualite=request.user.mutualite)
        if AffilieExclusion.objects.filter(affilie=affilie, medicament=item).exists():
            messages.warning(request, "Ce médicament est déjà exclu pour cet affilié.")
            return redirect('affilies:detail', pk=affilie_pk)
        AffilieExclusion.objects.create(
            mutualite=request.user.mutualite,
            affilie=affilie, medicament=item,
            motif=motif, cree_par=request.user,
        )
        messages.success(request, f"Médicament « {item.nom} » exclu pour {affilie.nom_complet}.")
    elif type_excl == 'service':
        from services_sante.models import ServiceSante
        item = get_object_or_404(ServiceSante, pk=item_pk, mutualite=request.user.mutualite)
        if AffilieExclusion.objects.filter(affilie=affilie, service=item).exists():
            messages.warning(request, "Ce service est déjà exclu pour cet affilié.")
            return redirect('affilies:detail', pk=affilie_pk)
        AffilieExclusion.objects.create(
            mutualite=request.user.mutualite,
            affilie=affilie, service=item,
            motif=motif, cree_par=request.user,
        )
        messages.success(request, f"Service « {item.nom} » exclu pour {affilie.nom_complet}.")

    return redirect('affilies:detail', pk=affilie_pk)


@login_required
@require_POST
def exclusion_supprimer(request, pk):
    from affilies.models import AffilieExclusion
    exclusion = get_object_or_404(AffilieExclusion, pk=pk, mutualite=request.user.mutualite)
    affilie_pk = exclusion.affilie.pk
    label = str(exclusion)
    exclusion.delete()
    messages.success(request, f"Exclusion supprimée : {label}")
    return redirect('affilies:detail', pk=affilie_pk)


@login_required
def exclusion_liste_ajax(request, affilie_pk):
    from affilies.models import AffilieExclusion
    affilie = get_object_or_404(Affilie, pk=affilie_pk, mutualite=request.user.mutualite)
    exclusions = AffilieExclusion.objects.filter(affilie=affilie).select_related('medicament', 'service')
    data = []
    for e in exclusions:
        item = {}
        if e.medicament:
            item = {'type': 'medicament', 'nom': e.medicament.nom, 'pk': e.medicament.pk}
        elif e.service:
            item = {'type': 'service', 'nom': e.service.nom, 'pk': e.service.pk}
        data.append({
            'id': e.pk,
            'item': item,
            'motif': e.motif,
            'date': e.date_creation.strftime('%d/%m/%Y'),
        })
    return JsonResponse({'success': True, 'data': data})


@login_required
def recherche_ajax(request):
    q = request.GET.get('q', '').strip()
    if len(q) < 1:
        return JsonResponse([], safe=False)
    affilies = Affilie.objects.filter(
        mutualite=request.user.mutualite, supprime=False
    ).filter(
        Q(nom__icontains=q) | Q(prenom__icontains=q) | Q(numero_affiliation__icontains=q)
    ).values('pk', 'numero_affiliation', 'nom', 'prenom')[:10]
    results = [{
        'id': a['pk'],
        'text': f"{a['numero_affiliation']} — {a['nom']} {a['prenom']}".strip()
    } for a in affilies]
    return JsonResponse(results, safe=False)