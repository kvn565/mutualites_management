from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.views.generic import TemplateView
from django.contrib.auth import get_user_model, authenticate, login
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.db import transaction
from django.contrib.auth.decorators import login_required
from django.utils import timezone
from datetime import timedelta
from django.core.paginator import Paginator
from django.db.models import Q, Count
from django.contrib.auth.hashers import make_password

import json
from django.db.models import Count, Sum
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core.models import JournalAudit, Mutualite, CleActivation, AuditCle, HistoriqueConnexion, ConfigurationApp, Notification, Utilisateur
from groupe.models import Groupe
from factures.models import Facture, LigneFacture
from affilies.models import Affilie
from ayants_droit.models import AyantDroit
from cotisations.models import Cotisation
from services_sante.models import ServiceSante
from origines.models import Origine
from partenaires.models import Partenaire, Convention, UtilisateurPartenaire, RelationPartenaireMutualite
from pharmacie.models import Medicament, CategorieMedicament
from partenaires.forms import PartenaireForm
from .forms import MutualiteForm, SuperAdminUtilisateurForm, ConfigurationAppForm, SuperAdminPartenaireForm, SuperAdminUtilisateurPartenaireForm, SuperAdminConventionForm

User = get_user_model()


class SuperAdminRequiredMixin(UserPassesTestMixin):
    def test_func(self):
        return self.request.user.is_authenticated and self.request.user.est_super_admin

    def handle_no_permission(self):
        messages.error(self.request, "Accès réservé aux super administrateurs.")
        return redirect('dashboard:index')

def superadmin_required(view_func):
    @login_required
    def wrapper(request, *args, **kwargs):
        if not (request.user.is_authenticated and request.user.est_super_admin):
            messages.error(request, "Accès réservé à l'administrateur.")
            return redirect('dashboard:index')
        return view_func(request, *args, **kwargs)
    return wrapper


class DashboardView(SuperAdminRequiredMixin, TemplateView):
    template_name = 'superadmin/dashboard.html'

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        now = timezone.now()

        total_membres = Affilie.objects.filter(supprime=False).count()
        total_factures = Facture.objects.count()

        ctx['stats'] = {
            'total_mutualites': Mutualite.objects.count(),
            'mutualites_essai': Mutualite.objects.filter(statut='ESSAI').count(),
            'mutualites_actives': Mutualite.objects.filter(statut='ACTIF').count(),
            'mutualites_suspendues': Mutualite.objects.filter(statut='SUSPENDU').count(),
            'mutualites_expirees': Mutualite.objects.filter(statut='EXPIRE').count(),
            'mutualites_bloquees': Mutualite.objects.filter(statut='BLOQUE').count(),
            'total_utilisateurs': User.objects.count(),
            'total_membres': total_membres,
            'total_factures': total_factures,
        }

        mutualites_qs = Mutualite.objects.annotate(
            nb_membres=Count('affilies', filter=Q(affilies__supprime=False)),
            nb_factures=Count('factures'),
        ).order_by('nom')

        ctx['chart_mutualites'] = json.dumps([
            {
                'nom': m.nom,
                'nb_membres': getattr(m, 'nb_membres', 0),
                'nb_factures': getattr(m, 'nb_factures', 0),
            }
            for m in mutualites_qs
        ])

        ctx['mutualites_liste'] = Mutualite.objects.all().order_by('nom')

        ctx['derniers_audits'] = JournalAudit.objects.filter(
            module='superadmin',
        ).order_by('-date_action')[:10]

        seuil_expiration = now + timedelta(days=3)
        ctx['essai_expirant'] = Mutualite.objects.filter(
            statut='ESSAI',
            date_fin_essai__lte=seuil_expiration,
        )

        ctx['mutualites_sans_licence'] = Mutualite.objects.filter(
            statut__in=['EN_ATTENTE', 'EXPIRE'],
        )

        seuil_7j = now + timedelta(days=7)
        ctx['expiration_proche'] = Mutualite.objects.filter(
            statut__in=['ACTIF', 'ESSAI'],
            date_expiration__lte=seuil_7j,
            date_expiration__gte=now,
        )

        return ctx


@superadmin_required
def mutualites_liste(request):
    q = request.GET.get('q', '')
    statut = request.GET.get('statut', '')

    mutualites = Mutualite.objects.all()

    if q:
        mutualites = mutualites.filter(
            Q(nom__icontains=q) | Q(nif__icontains=q) |
            Q(telephone__icontains=q) | Q(nom_responsable__icontains=q) |
            Q(code_mutualite__icontains=q)
        )

    if statut:
        mutualites = mutualites.filter(statut=statut.upper())

    paginator = Paginator(mutualites, 10)
    page_number = request.GET.get('page')
    mutualites_page = paginator.get_page(page_number)

    return render(request, 'superadmin/mutualites_liste.html', {
        'mutualites': mutualites_page,
        'q': q,
        'statut': statut,
        'paginator': paginator,
    })


@superadmin_required
def mutualite_creer(request):
    if request.method == 'POST':
        form = MutualiteForm(request.POST)
        if form.is_valid():
            with transaction.atomic():
                mutualite = form.save(commit=False)
                mutualite.created_by = request.user
                mutualite.save()

                admin_username = form.cleaned_data.get('username_admin') or f"admin_{mutualite.code_mutualite.lower()}"
                admin_email = form.cleaned_data.get('email_admin') or mutualite.email
                admin_password = form.cleaned_data.get('password_admin') or User.objects.make_random_password()

                admin_user = User.objects.create_user(
                    username=admin_username,
                    email=admin_email,
                    password=admin_password,
                    first_name=mutualite.nom_responsable.split()[0] if mutualite.nom_responsable else '',
                    last_name=' '.join(mutualite.nom_responsable.split()[1:]) if mutualite.nom_responsable else '',
                    role='ADMIN_MUTUALITE',
                    mutualite=mutualite,
                    telephone=mutualite.telephone,
                    actif=True,
                )

                mutualite.username_admin = admin_username
                mutualite.email_admin = admin_email
                mutualite.activer_essai(cree_par=request.user)

                JournalAudit.objects.create(
                    mutualite=mutualite,
                    utilisateur=request.user,
                    action='CREATION',
                    module='superadmin',
                    description=f"Mutualité '{mutualite.nom}' enregistrée par {request.user.username}",
                    objet_type='Mutualite',
                    objet_id=mutualite.pk,
                    adresse_ip=request.META.get('REMOTE_ADDR', ''),
                )

                messages.success(
                    request,
                    f"Mutualité « {mutualite.nom} » enregistrée. "
                    f"Administrateur créé : {admin_username} / mot de passe : {admin_password}. "
                    f"Essai gratuit de 14 jours activé !"
                )
                return redirect('superadmin:mutualite_detail', pk=mutualite.pk)
    else:
        form = MutualiteForm()

    return render(request, 'superadmin/mutualite_form.html', {
        'form': form,
        'titre': "Enregistrer une nouvelle mutualité",
    })


@superadmin_required
def mutualite_modifier(request, pk):
    mutualite = get_object_or_404(Mutualite, pk=pk)
    if request.method == 'POST':
        form = MutualiteForm(request.POST, instance=mutualite)
        if form.is_valid():
            form.save()
            JournalAudit.objects.create(
                mutualite=mutualite,
                utilisateur=request.user,
                action='MODIFICATION',
                module='superadmin',
                description=f"Mutualité '{mutualite.nom}' modifiée par {request.user.username}",
                objet_type='Mutualite',
                objet_id=mutualite.pk,
                adresse_ip=request.META.get('REMOTE_ADDR', ''),
            )
            messages.success(request, f"Mutualité « {mutualite.nom} » modifiée.")
            return redirect('superadmin:mutualite_detail', pk=pk)
    else:
        form = MutualiteForm(instance=mutualite)

    return render(request, 'superadmin/mutualite_form.html', {
        'form': form,
        'titre': 'Modifier mutualité',
        'mutualite': mutualite,
    })


@superadmin_required
def mutualite_detail(request, pk):
    mutualite = get_object_or_404(Mutualite, pk=pk)
    # Get partners via FK AND via RelationPartenaireMutualite
    from partenaires.models import RelationPartenaireMutualite
    partenaires_fk = Partenaire.objects.filter(mutualite=mutualite)
    relations = RelationPartenaireMutualite.objects.filter(mutualite=mutualite).select_related('partenaire')
    partenaires_rel = [r.partenaire for r in relations if r.partenaire not in partenaires_fk]
    partenaires = list(partenaires_fk) + partenaires_rel
    partenaires.sort(key=lambda p: p.date_creation, reverse=True)
    
    journaux = JournalAudit.objects.filter(mutualite=mutualite).order_by('-date_action')[:20]

    return render(request, 'superadmin/mutualite_detail.html', {
        'mutualite': mutualite,
        'partenaires': partenaires,
        'journaux': journaux,
    })


@superadmin_required
def mutualite_suspendre(request, pk):
    mutualite = get_object_or_404(Mutualite, pk=pk)
    if mutualite.statut not in ('ACTIF', 'ESSAI'):
        messages.warning(request, f"La mutualité « {mutualite.nom} » ne peut pas être suspendue (statut : {mutualite.get_statut_display()}).")
        return redirect('superadmin:mutualite_detail', pk=pk)

    mutualite.statut = 'SUSPENDU'
    mutualite.save()

    JournalAudit.objects.create(
        mutualite=mutualite,
        utilisateur=request.user,
        action='MODIFICATION',
        module='superadmin',
        description=f"Mutualité '{mutualite.nom}' suspendue par {request.user.username}",
        objet_type='Mutualite',
        objet_id=mutualite.pk,
        adresse_ip=request.META.get('REMOTE_ADDR', ''),
    )
    messages.warning(request, f"Mutualité « {mutualite.nom} » suspendue.")
    return redirect('superadmin:mutualite_detail', pk=pk)


@superadmin_required
def mutualite_reactiver(request, pk):
    mutualite = get_object_or_404(Mutualite, pk=pk)
    if mutualite.statut != 'SUSPENDU':
        messages.info(request, f"La mutualité « {mutualite.nom} » n'est pas suspendue.")
        return redirect('superadmin:mutualite_detail', pk=pk)

    if mutualite.essai_gratuit:
        mutualite.activer_essai()
    else:
        mutualite.statut = 'ACTIF'
        mutualite.save()

    JournalAudit.objects.create(
        mutualite=mutualite,
        utilisateur=request.user,
        action='MODIFICATION',
        module='superadmin',
        description=f"Mutualité '{mutualite.nom}' réactivée par {request.user.username}",
        objet_type='Mutualite',
        objet_id=mutualite.pk,
        adresse_ip=request.META.get('REMOTE_ADDR', ''),
    )
    messages.success(request, f"Mutualité « {mutualite.nom} » réactivée.")
    return redirect('superadmin:mutualite_detail', pk=pk)


@superadmin_required
def mutualite_renouveler_licence(request, pk):
    mutualite = get_object_or_404(Mutualite, pk=pk)
    if request.method == 'POST':
        type_abonnement = request.POST.get('type_abonnement', 'MENSUEL')
        duree_jours = int(request.POST.get('duree_jours', 30))

        mutualite.activer_licence(type_abonnement, duree_jours)

        JournalAudit.objects.create(
            mutualite=mutualite,
            utilisateur=request.user,
            action='MODIFICATION',
            module='superadmin',
            description=f"Licence renouvelée pour '{mutualite.nom}' : {type_abonnement}, {duree_jours} jours, par {request.user.username}",
            objet_type='Mutualite',
            objet_id=mutualite.pk,
            adresse_ip=request.META.get('REMOTE_ADDR', ''),
        )
        messages.success(request, f"Licence renouvelée pour « {mutualite.nom} » jusqu'au {mutualite.date_expiration.strftime('%d/%m/%Y')}.")
        return redirect('superadmin:mutualite_detail', pk=pk)

    return render(request, 'superadmin/renouveler_licence.html', {
        'mutualite': mutualite,
    })


@superadmin_required
def mutualite_essai_prolonger(request, pk):
    mutualite = get_object_or_404(Mutualite, pk=pk)
    if request.method == 'POST':
        jours = int(request.POST.get('jours', 7))
        from datetime import timedelta
        if mutualite.date_fin_essai:
            mutualite.date_fin_essai += timedelta(days=jours)
        else:
            mutualite.date_fin_essai = timezone.now() + timedelta(days=jours)
        mutualite.date_expiration = mutualite.date_fin_essai
        mutualite.statut = 'ESSAI'
        mutualite.save()

        JournalAudit.objects.create(
            mutualite=mutualite,
            utilisateur=request.user,
            action='MODIFICATION',
            module='superadmin',
            description=f"Essai prolongé de {jours} jours pour '{mutualite.nom}' par {request.user.username}",
            objet_type='Mutualite',
            objet_id=mutualite.pk,
            adresse_ip=request.META.get('REMOTE_ADDR', ''),
        )
        messages.success(request, f"Essai prolongé de {jours} jours pour « {mutualite.nom} ».")
        return redirect('superadmin:mutualite_detail', pk=pk)

    return render(request, 'superadmin/prolonger_essai.html', {
        'mutualite': mutualite,
    })


@superadmin_required
def utilisateurs_liste(request):
    q = request.GET.get('q', '')
    role = request.GET.get('role', '')
    mutualite_id = request.GET.get('mutualite', '')

    utilisateurs = User.objects.all().order_by('-date_creation')

    if q:
        utilisateurs = utilisateurs.filter(
            Q(username__icontains=q) |
            Q(email__icontains=q) |
            Q(nom__icontains=q) |
            Q(prenom__icontains=q) |
            Q(telephone__icontains=q)
        )
    if role:
        utilisateurs = utilisateurs.filter(role=role)
    if mutualite_id:
        utilisateurs = utilisateurs.filter(mutualite_id=mutualite_id)

    paginator = Paginator(utilisateurs, 10)
    page_number = request.GET.get('page')
    utilisateurs_page = paginator.get_page(page_number)

    mutualites = Mutualite.objects.all().order_by('nom')

    return render(request, 'superadmin/utilisateurs_liste.html', {
        'utilisateurs': utilisateurs_page,
        'paginator': paginator,
        'q': q,
        'role': role,
        'mutualite_id': mutualite_id,
        'mutualites': mutualites,
    })


@superadmin_required
def utilisateur_creer(request):
    if request.method == 'POST':
        form = SuperAdminUtilisateurForm(request.POST)
        if form.is_valid():
            user = form.save()
            JournalAudit.objects.create(
                mutualite=user.mutualite,
                utilisateur=request.user,
                action='CREATION',
                module='superadmin',
                description=f"Utilisateur '{user.username}' créé par {request.user.username}",
                objet_type='Utilisateur',
                objet_id=user.pk,
                adresse_ip=request.META.get('REMOTE_ADDR', ''),
            )
            messages.success(request, f"Utilisateur « {user.nom_complet or user.username} » créé avec succès.")
            return redirect('superadmin:utilisateurs_liste')
    else:
        form = SuperAdminUtilisateurForm()

    return render(request, 'superadmin/utilisateur_form.html', {
        'form': form,
        'titre': 'Créer un utilisateur',
        'is_create': True,
    })


@superadmin_required
def utilisateur_modifier(request, pk):
    user = get_object_or_404(User, pk=pk)
    if request.method == 'POST':
        form = SuperAdminUtilisateurForm(request.POST, instance=user)
        if form.is_valid():
            user = form.save()
            JournalAudit.objects.create(
                mutualite=user.mutualite,
                utilisateur=request.user,
                action='MODIFICATION',
                module='superadmin',
                description=f"Utilisateur '{user.username}' modifié par {request.user.username}",
                objet_type='Utilisateur',
                objet_id=user.pk,
                adresse_ip=request.META.get('REMOTE_ADDR', ''),
            )
            messages.success(request, f"Utilisateur « {user.nom_complet or user.username} » modifié.")
            return redirect('superadmin:utilisateurs_liste')
    else:
        form = SuperAdminUtilisateurForm(instance=user)

    return render(request, 'superadmin/utilisateur_form.html', {
        'form': form,
        'titre': "Modifier l'utilisateur",
        'user': user,
    })


@superadmin_required
def utilisateur_supprimer(request, pk):
    user = get_object_or_404(User, pk=pk)
    if user == request.user:
        messages.error(request, 'Vous ne pouvez pas supprimer votre propre compte.')
        return redirect('superadmin:utilisateurs_liste')
    if user.est_super_admin:
        messages.error(request, 'Vous ne pouvez pas supprimer un super administrateur.')
        return redirect('superadmin:utilisateurs_liste')

    if request.method == 'POST':
        nom = user.nom_complet or user.username
        user.delete()
        JournalAudit.objects.create(
            mutualite=user.mutualite,
            utilisateur=request.user,
            action='SUPPRESSION',
            module='superadmin',
            description=f"Utilisateur '{nom}' supprimé par {request.user.username}",
            objet_type='Utilisateur',
            objet_id=pk,
            adresse_ip=request.META.get('REMOTE_ADDR', ''),
        )
        messages.success(request, f"Utilisateur « {nom} » supprimé.")
        return redirect('superadmin:utilisateurs_liste')

    return render(request, 'superadmin/utilisateur_confirm_supprimer.html', {
        'user': user,
    })


@superadmin_required
def cles_liste(request):
    q = request.GET.get('q', '')
    statut = request.GET.get('statut', '')
    type_plan = request.GET.get('type_plan', '')

    cles = CleActivation.objects.all().select_related('mutualite')

    if q:
        cles = cles.filter(
            Q(cle_visible__icontains=q) |
            Q(mutualite__nom__icontains=q) |
            Q(mutualite__nif__icontains=q)
        )
    if statut:
        cles = cles.filter(statut=statut)
    if type_plan:
        cles = cles.filter(type_plan=type_plan)

    paginator = Paginator(cles, 10)
    page_number = request.GET.get('page')
    cles_page = paginator.get_page(page_number)

    return render(request, 'superadmin/cles_liste.html', {
        'cles': cles_page,
        'paginator': paginator,
        'q': q,
        'statut': statut,
        'type_plan': type_plan,
    })


@superadmin_required
def cle_generer(request):
    mutualites = Mutualite.objects.all().order_by('nom')

    if request.method == 'POST':
        mutualite_id = request.POST.get('mutualite')
        type_plan = request.POST.get('type_plan')
        date_debut = request.POST.get('date_debut')
        date_fin = request.POST.get('date_fin')

        if not mutualite_id or not type_plan:
            messages.error(request, 'Veuillez sélectionner une mutualité et un type de plan.')
            return redirect('superadmin:cles_generer')

        mutualite = get_object_or_404(Mutualite, pk=mutualite_id)

        try:
            from datetime import datetime

            if type_plan == 'PERSONNALISE':
                if not date_debut or not date_fin:
                    messages.error(request, 'Veuillez spécifier les dates de début et fin pour un plan personnalisé.')
                    return redirect('superadmin:cles_generer')
                dt_debut = datetime.strptime(date_debut, '%Y-%m-%d').replace(tzinfo=timezone.get_current_timezone())
                dt_fin = datetime.strptime(date_fin, '%Y-%m-%d').replace(tzinfo=timezone.get_current_timezone())
            else:
                dt_debut = timezone.now()
                plan_to_duree = {
                    'ESSAI': 14, '1MOIS': 30, '6MOIS': 182,
                    '12MOIS': 365, '24MOIS': 730,
                }
                duree_jours = plan_to_duree.get(type_plan, 30)
                dt_debut = timezone.now()
                dt_fin = dt_debut + timedelta(days=duree_jours)

            cle = CleActivation.objects.create(
                mutualite=mutualite,
                type_plan=type_plan,
                date_debut=dt_debut,
                date_fin=dt_fin,
                active=True,
                utilisee=False,
                statut='DISPONIBLE',
                cree_par=request.user.username,
                notes=f"Générée par {request.user.username}",
            )

            AuditCle.objects.create(
                cle=cle,
                mutualite=mutualite,
                action='CREEE',
                message=f"Clé {cle.cle_visible} générée pour {mutualite.nom} ({cle.label_plan})",
                ip_address=request.META.get('REMOTE_ADDR', ''),
            )

            messages.success(
                request,
                f"Clé générée avec succès pour « {mutualite.nom} ».<br>"
                f"<strong>Clé :</strong> <code>{cle.cle_visible}</code><br>"
                f"<strong>Plan :</strong> {cle.label_plan}<br>"
                f"<strong>Validité :</strong> du {dt_debut.strftime('%d/%m/%Y')} au {dt_fin.strftime('%d/%m/%Y')}"
            )
            return redirect('superadmin:cles_liste')

        except Exception as e:
            messages.error(request, f"Erreur lors de la génération : {e}")
            return redirect('superadmin:cles_generer')

    return render(request, 'superadmin/cle_generer.html', {
        'mutualites': mutualites,
    })


@superadmin_required
def cle_activer(request, pk):
    cle = get_object_or_404(CleActivation, pk=pk)
    mutualite = cle.mutualite

    if cle.statut != 'DISPONIBLE':
        messages.warning(request, f"Cette clé a déjà le statut « {cle.get_statut_display()} ».")
        return redirect('superadmin:cles_liste')

    cle.utilisee = True
    cle.date_utilisation = timezone.now()
    cle.save()

    mutualite.cle_activation = cle.cle_visible
    mutualite.date_activation = cle.date_debut
    mutualite.date_expiration = cle.date_fin
    mutualite.type_abonnement = cle.type_plan
    mutualite.essai_gratuit = (cle.type_plan == 'ESSAI')
    mutualite.statut = 'ACTIF'
    mutualite.save()

    AuditCle.objects.create(
        cle=cle,
        mutualite=mutualite,
        action='ACTIVEE',
        message=f"Clé {cle.cle_visible} activée pour {mutualite.nom}",
        ip_address=request.META.get('REMOTE_ADDR', ''),
    )

    messages.success(request, f"Clé activée pour « {mutualite.nom} ». Expire le {cle.date_fin.strftime('%d/%m/%Y')}.")
    return redirect('superadmin:cles_liste')


@superadmin_required
def cle_revoquer(request, pk):
    cle = get_object_or_404(CleActivation, pk=pk)
    if cle.statut == 'REVOQUEE':
        messages.warning(request, 'Cette clé est déjà révoquée.')
        return redirect('superadmin:cles_liste')

    if request.method == 'POST':
        motif = request.POST.get('motif', '')
        cle.active = False
        cle.motif_revocation = motif
        cle.save()

        AuditCle.objects.create(
            cle=cle,
            mutualite=cle.mutualite,
            action='REVOQUEE',
            message=f"Clé {cle.cle_visible} révoquée : {motif or 'Aucun motif'}",
            ip_address=request.META.get('REMOTE_ADDR', ''),
        )

        messages.success(request, f"Clé « {cle.cle_visible} » révoquée.")
        return redirect('superadmin:cles_liste')

    return render(request, 'superadmin/cle_revoquer.html', {'cle': cle})


@superadmin_required
def mutualite_factures(request, pk):
    mutualite = get_object_or_404(Mutualite, pk=pk)
    factures = Facture.objects.filter(mutualite=mutualite).select_related(
        'partenaire', 'affilie', 'ayant_droit'
    ).order_by('-date', '-pk')
    paginator = Paginator(factures, 10)

    page = request.GET.get('page')
    factures_page = paginator.get_page(page)

    from django.db.models import Sum
    stats = Facture.objects.filter(mutualite=mutualite).aggregate(
        total=Sum('montant_total'),
        part_mutualite=Sum('montant_mutualite'),
        part_patient=Sum('montant_ticket_moderateur'),
        nb_brouillon=Count('pk', filter=Q(statut='BROUILLON')),
        nb_soumis=Count('pk', filter=Q(statut='SOUMIS')),
        nb_valide=Count('pk', filter=Q(statut='VALIDE')),
        nb_paye=Count('pk', filter=Q(statut='PAYE')),
    )

    return render(request, 'superadmin/mutualite_factures.html', {
        'mutualite': mutualite,
        'factures': factures_page,
        'paginator': paginator,
        'stats': stats,
    })


@superadmin_required
def mutualite_membres(request, pk):
    mutualite = get_object_or_404(Mutualite, pk=pk)
    affilies = Affilie.objects.filter(mutualite=mutualite, supprime=False).order_by('-date_adhesion')

    paginator = Paginator(affilies, 10)
    page = request.GET.get('page')
    affilies_page = paginator.get_page(page)

    stats = Affilie.objects.filter(mutualite=mutualite, supprime=False).aggregate(
        total=Count('pk'),
        actifs=Count('pk', filter=Q(statut='ACTIF')),
        en_attente=Count('pk', filter=Q(statut='EN_ATTENTE')),
        suspendus=Count('pk', filter=Q(statut='SUSPENDU')),
        radios=Count('pk', filter=Q(statut='RADIE')),
    )

    return render(request, 'superadmin/mutualite_membres.html', {
        'mutualite': mutualite,
        'affilies': affilies_page,
        'paginator': paginator,
        'stats': stats,
    })


@superadmin_required
def suivi_factures(request):
    mutualites = Mutualite.objects.annotate(
        nb_factures=Count('factures'),
    ).order_by('nom')

    return render(request, 'superadmin/suivi_factures.html', {
        'mutualites': mutualites,
    })


@superadmin_required
def suivi_membres(request):
    mutualites = Mutualite.objects.annotate(
        nb_affilies=Count('affilies', filter=Q(affilies__supprime=False)),
    ).order_by('nom')

    return render(request, 'superadmin/membres_liste.html', {
        'mutualites': mutualites,
    })


@superadmin_required
def audit_connexions(request):
    q = request.GET.get('q', '')
    mutualite_id = request.GET.get('mutualite', '')

    connexions = HistoriqueConnexion.objects.select_related(
        'utilisateur', 'mutualite'
    ).order_by('-date_connexion')

    if q:
        connexions = connexions.filter(
            Q(utilisateur__username__icontains=q) |
            Q(utilisateur__nom__icontains=q) |
            Q(utilisateur__prenom__icontains=q) |
            Q(adresse_ip__icontains=q) |
            Q(user_agent__icontains=q)
        )
    if mutualite_id:
        connexions = connexions.filter(mutualite_id=mutualite_id)

    paginator = Paginator(connexions, 10)
    page = request.GET.get('page')
    connexions_page = paginator.get_page(page)

    mutualites = Mutualite.objects.all().order_by('nom')

    return render(request, 'superadmin/audit_connexions.html', {
        'connexions': connexions_page,
        'paginator': paginator,
        'q': q,
        'mutualite_id': mutualite_id,
        'mutualites': mutualites,
    })

from django.http import HttpResponse, JsonResponse

@superadmin_required
def reinitialiser_selection(request):
    mutualites = Mutualite.objects.all().order_by('nom')
    selected = None
    pk = request.GET.get('mutualite')
    if pk:
        try:
            selected = Mutualite.objects.get(pk=pk)
        except Mutualite.DoesNotExist:
            pass
    return render(request, 'superadmin/reinitialiser_selection.html', {
        'mutualites': mutualites,
        'selected': selected,
    })

@superadmin_required
def cle_irrevoquer(request, pk):
    return HttpResponse("Fonctionnalité en cours de développement.")

@superadmin_required
def mutualite_reinitialiser(request, pk):
    mutualite = get_object_or_404(Mutualite, pk=pk)

    # Statistiques des données opérationnelles
    from affilies.models import Affilie
    from ayants_droit.models import AyantDroit
    from cotisations.models import Cotisation
    from factures.models import Facture
    from services_sante.models import ServiceSante
    from origines.models import Origine

    stats = {
        'affilies': Affilie.objects.filter(mutualite=mutualite, supprime=False).count(),
        'ayants_droit': AyantDroit.objects.filter(mutualite=mutualite, supprime=False).count(),
        'cotisations': Cotisation.objects.filter(mutualite=mutualite).count(),
        'factures': Facture.objects.filter(mutualite=mutualite).count(),
        'services': ServiceSante.objects.filter(mutualite=mutualite).count(),
        'partenaires': Partenaire.objects.filter(mutualite=mutualite).count(),
        'origines': Origine.objects.filter(mutualite=mutualite).count(),
    }
    total = sum(stats.values())

    if request.method == 'POST':
        action = request.POST.get('action')

        if action == 'backup':
            data = {'mutualite': {'nom': mutualite.nom, 'code': mutualite.code_mutualite}}
            data['affilies'] = list(Affilie.objects.filter(mutualite=mutualite).values())
            data['ayants_droit'] = list(AyantDroit.objects.filter(mutualite=mutualite).values())
            data['cotisations'] = list(Cotisation.objects.filter(mutualite=mutualite).values())
            data['factures'] = list(Facture.objects.filter(mutualite=mutualite).values())
            data['services'] = list(ServiceSante.objects.filter(mutualite=mutualite).values())
            data['partenaires'] = list(Partenaire.objects.filter(mutualite=mutualite).values())
            data['origines'] = list(Origine.objects.filter(mutualite=mutualite).values())

            response = HttpResponse(
                json.dumps(data, indent=2, default=str),
                content_type='application/json',
            )
            response['Content-Disposition'] = f'attachment; filename="backup_{mutualite.code_mutualite}_{timezone.now().strftime("%Y%m%d_%H%M%S")}.json"'
            return response

        if action == 'reinitialiser':
            confirm = request.POST.get('confirm', '').strip()
            if confirm != mutualite.nom:
                messages.error(request, "Le nom saisi ne correspond pas.")
                return redirect('superadmin:mutualite_reinitialiser', pk=pk)

            with transaction.atomic():
                Origine.objects.filter(mutualite=mutualite).delete()
                ServiceSante.objects.filter(mutualite=mutualite).delete()
                Facture.objects.filter(mutualite=mutualite).delete()
                Cotisation.objects.filter(mutualite=mutualite).delete()
                AyantDroit.objects.filter(mutualite=mutualite).delete()
                Affilie.objects.filter(mutualite=mutualite).delete()
                Partenaire.objects.filter(mutualite=mutualite).delete()

            JournalAudit.objects.create(
                mutualite=mutualite, utilisateur=request.user,
                action='REINITIALISATION', module='superadmin',
                description=f"Réinitialisation des données de {mutualite.nom}",
                objet_type='Mutualite', objet_id=mutualite.pk,
                adresse_ip=request.META.get('REMOTE_ADDR', ''),
            )
            messages.success(request, f"Toutes les données de {mutualite.nom} ont été supprimées.")
            return redirect('superadmin:mutualite_reinitialiser', pk=pk)

    return render(request, 'superadmin/mutualite_reinitialiser.html', {
        'mutualite': mutualite,
        'stats': stats,
        'total': total,
    })

def _excel_style_header(ws, headers):
    bold = Font(bold=True, color='FFFFFF', size=11)
    fill = PatternFill(start_color='0a5c36', end_color='0a5c36', fill_type='solid')
    align = Alignment(horizontal='center', vertical='center')
    thin = Side(style='thin', color='c9a84c')
    border = Border(top=thin, left=thin, right=thin)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = bold
        cell.fill = fill
        cell.alignment = align
        cell.border = border

def _excel_write_rows(ws, rows):
    thin = Side(style='thin', color='dde3ea')
    border = Border(bottom=thin)
    for r, row_data in enumerate(rows, 2):
        for c, val in enumerate(row_data, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.border = border

def _build_excel_mutualite(mutualite):
    wb = openpyxl.Workbook()
    del wb['Sheet']
    m = mutualite

    def _add_sheet(name, headers, rows):
        ws = wb.create_sheet(title=name[:31])
        _excel_style_header(ws, headers)
        _excel_write_rows(ws, rows)
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = max(12, min(40, len(str(headers[col - 1])) + 4))

    def _date(val):
        return val.strftime('%d/%m/%Y') if val else ''

    def _datetime(val):
        return val.strftime('%d/%m/%Y %H:%M') if val else ''

    # 1. Mutualite
    _add_sheet('Mutualité',
        ['NIF', 'Nom', 'Sigle', 'Type', 'Statut', 'Responsable', 'Téléphone', 'Email',
         'Province', 'Commune', 'Adresse', 'Code', 'Max utilisateurs', 'Créé le'],
        [[m.nif, m.nom, m.sigle, m.type_mutualite, m.statut, m.nom_responsable,
          m.telephone, m.email, m.province, m.commune, m.adresse, m.code_mutualite,
          m.max_utilisateurs, _date(m.date_creation)]])

    # 2. Utilisateurs
    rows = list(Utilisateur.objects.filter(mutualite=mutualite).values_list(
        'username', 'email', 'nom', 'postnom', 'prenom', 'role', 'telephone', 'actif'))
    _add_sheet('Utilisateurs',
        ['Username', 'Email', 'Nom', 'Postnom', 'Prénom', 'Rôle', 'Téléphone', 'Actif'], rows)

    # 3. Groupes
    rows = list(Groupe.objects.filter(mutualite=mutualite).values_list('nom', 'description'))
    _add_sheet('Groupes', ['Nom', 'Description'], rows)

    # 4. Affiliés
    qs = Affilie.objects.filter(mutualite=mutualite).values_list(
        'numero_affiliation', 'nom', 'prenom', 'sexe', 'telephone', 'email',
        'profession', 'statut', 'date_adhesion')
    rows = [[a[0], a[1], a[2], a[3], a[4], a[5], a[6], a[7], _date(a[8])] for a in qs]
    _add_sheet('Affiliés',
        ['N° Affiliation', 'Nom', 'Prénom', 'Sexe', 'Téléphone', 'Email',
         'Profession', 'Statut', 'Date adhésion'], rows)

    # 5. Ayants droit
    qs = AyantDroit.objects.filter(mutualite=mutualite).values_list(
        'nom', 'prenom', 'sexe', 'date_naissance', 'lien_parente', 'affilie__numero_affiliation')
    rows = [[a[0], a[1], a[2], _date(a[3]), a[4], a[5]] for a in qs]
    _add_sheet('Ayants droit',
        ['Nom', 'Prénom', 'Sexe', 'Date naissance', 'Lien parenté', 'Affilié'], rows)

    # 6. Cotisations
    qs = Cotisation.objects.filter(mutualite=mutualite).values_list(
        'affilie__numero_affiliation', 'montant', 'mois', 'annee', 'date_paiement', 'statut', 'reference')
    rows = [[c[0], c[1], c[2], c[3], _date(c[4]), c[5], c[6]] for c in qs]
    _add_sheet('Cotisations',
        ['Affilié', 'Montant', 'Mois', 'Année', 'Date paiement', 'Statut', 'Référence'], rows)

    # 7. Factures
    qs = Facture.objects.filter(mutualite=mutualite).values_list(
        'numero', 'affilie__numero_affiliation', 'date', 'montant_total', 'statut')
    rows = [[f[0], f[1], _date(f[2]), f[3], f[4]] for f in qs]
    _add_sheet('Factures', ['N° Facture', 'Affilié', 'Date', 'Total', 'Statut'], rows)

    # 8. Lignes facture
    qs = LigneFacture.objects.filter(mutualite=mutualite).values_list(
        'facture__numero', 'designation', 'quantite', 'prix_unitaire')
    rows = [[l[0], l[1], l[2], l[3], l[2] * l[3]] for l in qs]
    _add_sheet('Lignes facture',
        ['N° Facture', 'Désignation', 'Qté', 'Prix unitaire', 'Montant'], rows)

    # 9. Services santé
    rows = list(ServiceSante.objects.filter(mutualite=mutualite).values_list(
        'code', 'nom', 'categorie', 'plafond', 'actif'))
    _add_sheet('Services santé', ['Code', 'Nom', 'Catégorie', 'Plafond', 'Actif'], rows)

    # 10. Partenaires
    rows = list(Partenaire.objects.filter(mutualite=mutualite).values_list(
        'code_partenaire', 'nom_etablissement', 'type_partenaire', 'telephone', 'email', 'statut'))
    _add_sheet('Partenaires',
        ['Code', 'Établissement', 'Type', 'Téléphone', 'Email', 'Statut'], rows)

    # 11. Conventions
    qs = Convention.objects.filter(mutualite=mutualite).values_list(
        'numero_convention', 'partenaire__nom_etablissement', 'date_debut', 'date_fin', 'taux_prise_en_charge', 'statut')
    rows = [[c[0], c[1], _date(c[2]), _date(c[3]), c[4], c[5]] for c in qs]
    _add_sheet('Conventions',
        ['N° Convention', 'Partenaire', 'Début', 'Fin', 'Taux prise en charge', 'Statut'], rows)

    # 12. Utilisateurs partenaires
    rows = list(UtilisateurPartenaire.objects.filter(mutualite=mutualite).values_list(
        'username', 'role', 'email', 'telephone', 'actif'))
    _add_sheet('Utilisateurs partenaires',
        ['Username', 'Rôle', 'Email', 'Téléphone', 'Actif'], rows)

    # 13. Origines
    rows = list(Origine.objects.filter(mutualite=mutualite).values_list('nom', 'description'))
    _add_sheet('Origines', ['Nom', 'Description'], rows)

    # 14. Médicaments
    rows = list(Medicament.objects.filter(mutualite=mutualite).values_list(
        'code', 'nom', 'forme', 'dosage', 'prix_reference', 'actif'))
    _add_sheet('Médicaments', ['Code', 'Nom', 'Forme', 'Dosage', 'Prix référence', 'Actif'], rows)

    # 15. Notifications
    qs = Notification.objects.filter(mutualite=mutualite).values_list(
        'type_notification', 'message', 'vue', 'date_creation')
    rows = [[n[0], n[1], 'Oui' if n[2] else 'Non', _datetime(n[3])] for n in qs]
    _add_sheet('Notifications', ['Type', 'Message', 'Vue', 'Date'], rows)

    # 16. Journal audit
    qs = JournalAudit.objects.filter(mutualite=mutualite).values_list(
        'action', 'module', 'description', 'utilisateur__username', 'date_action')
    rows = [[a[0], a[1], a[2], a[3], _datetime(a[4])] for a in qs]
    _add_sheet('Journal audit', ['Action', 'Module', 'Description', 'Utilisateur', 'Date'], rows)

    return wb


@superadmin_required
def export_mutualite_excel(request, pk):
    mutualite = get_object_or_404(Mutualite, pk=pk)
    wb = _build_excel_mutualite(mutualite)
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    code = mutualite.code_mutualite or f'mut{mutualite.pk}'
    response['Content-Disposition'] = f'attachment; filename="export_{code}_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
    wb.save(response)
    return response

@superadmin_required
def export_toutes_mutualites_excel(request):
    mutualites = Mutualite.objects.all().order_by('nom')
    wb = openpyxl.Workbook()
    del wb['Sheet']

    for mutualite in mutualites:
        wb_mut = _build_excel_mutualite(mutualite)
        for ws in wb_mut.worksheets:
            new_ws = wb.create_sheet(title=f'{mutualite.sigle or mutualite.nom[:4]}_{ws.title}'[:31])
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column, values_only=True):
                new_ws.append(row)
            _excel_style_header(new_ws, [cell.value for cell in new_ws[1]] if new_ws.max_row > 0 else [])
            for col in range(1, new_ws.max_column + 1):
                new_ws.column_dimensions[get_column_letter(col)].width = max(12, min(40, 16))

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="export_toutes_mutualites_{timezone.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
    wb.save(response)
    return response

def ajax_verifier_nif(request):
    """
    Vérifie un NIF lors de l'inscription d'une mutualité.

    1. Base locale d'abord : si le NIF correspond à une mutualité déjà créée,
       on retourne toutes ses informations (pré-remplissage) + le statut 'existe'.
    2. Sinon, on interroge l'API de l'OBR pour retrouver le nom légal.
    """
    nif = (request.GET.get('nif') or request.POST.get('nif') or '').strip()
    if not nif:
        return JsonResponse({"ok": False, "error": "NIF requis"})

    from core.models import Mutualite
    mutualite = Mutualite.objects.filter(nif=nif).first()
    if mutualite:
        return JsonResponse({
            "ok": True,
            "status": "existe",
            "entite": "mutualite",
            "nom_local": mutualite.nom,
            "nom": mutualite.nom,
            "sigle": mutualite.sigle,
            "province": mutualite.province,
            "commune": mutualite.commune,
            "quartier": mutualite.quartier,
            "adresse": mutualite.adresse,
            "telephone": mutualite.telephone,
            "email": mutualite.email,
        })

    from partenaires.models import Partenaire
    partenaire = Partenaire.all_objects.filter(nif=nif, is_deleted=False).first()
    if partenaire:
        return JsonResponse({
            "ok": True,
            "status": "existe",
            "entite": "partenaire",
            "nom_local": partenaire.nom_etablissement,
            "nom_etablissement": partenaire.nom_etablissement,
            "email": partenaire.email,
            "telephone": partenaire.telephone,
            "province": partenaire.province,
            "commune": partenaire.commune,
            "zone": partenaire.zone,
            "quartier": partenaire.quartier,
            "avenue": partenaire.avenue,
            "numero": partenaire.numero,
            "nom_responsable": partenaire.nom_responsable,
        })

    from core.services import verifier_nif_obr
    resultat = verifier_nif_obr(nif)
    if resultat.get('ok'):
        return JsonResponse({"ok": True, "status": "nouveau", "nom": resultat['nom']})
    return JsonResponse({"ok": False, "error": resultat.get('error', 'NIF non trouvé sur OBR')})

@superadmin_required
def configuration_app(request):
    config = ConfigurationApp.get_config()
    if request.method == 'POST':
        form = ConfigurationAppForm(request.POST, request.FILES, instance=config)
        if form.is_valid():
            form.save()
            messages.success(request, 'Configuration mise à jour avec succès.')
            return redirect('superadmin:configuration_app')
    else:
        form = ConfigurationAppForm(instance=config)

    return render(request, 'superadmin/configuration_app.html', {
        'form': form,
        'config': config,
    })

@superadmin_required
def partenaires_liste(request):
    qs = Partenaire.all_objects.filter(supprime=False).select_related('mutualite')

    q = request.GET.get('q', '')
    type_part = request.GET.get('type', '')
    statut = request.GET.get('statut', '')
    mutualite_id = request.GET.get('mutualite', '')

    if q:
        qs = qs.filter(
            Q(nom_etablissement__icontains=q) |
            Q(code_partenaire__icontains=q) |
            Q(nif__icontains=q)
        )
    if type_part:
        qs = qs.filter(type_partenaire=type_part)
    if statut:
        qs = qs.filter(statut=statut)
    if mutualite_id:
        qs = qs.filter(mutualite_id=mutualite_id)

    paginator = Paginator(qs, 15)
    page = request.GET.get('page', 1)
    partenaires = paginator.get_page(page)

    return render(request, 'superadmin/partenaires_liste.html', {
        'partenaires': partenaires,
        'q': q,
        'type': type_part,
        'statut': statut,
        'mutualite_id': mutualite_id,
        'type_choices': Partenaire.TYPE_CHOICES,
        'statut_choices': Partenaire.STATUT_CHOICES,
        'mutualites': Mutualite.objects.all().order_by('nom'),
    })

@superadmin_required
def partenaire_creer(request):
    if request.method == 'POST':
        form = SuperAdminPartenaireForm(request.POST, request.FILES)
        if form.is_valid():
            obj = form.save(commit=False)
            obj.created_by = request.user
            if not obj.code_partenaire:
                obj.code_partenaire = Partenaire.generer_code(obj.type_partenaire)
            obj.save()
            JournalAudit.objects.create(
                mutualite=obj.mutualite, utilisateur=request.user,
                action='CREATION', module='superadmin',
                description=f"Création partenaire : {obj.nom_etablissement}",
                objet_type='Partenaire', objet_id=obj.pk,
                adresse_ip=request.META.get('REMOTE_ADDR', ''),
            )
            messages.success(request, f"Partenaire {obj.nom_etablissement} créé.")
            return redirect('superadmin:partenaire_detail', pk=obj.pk)
    else:
        initial = {}
        t = request.GET.get('type', '')
        if t:
            initial['code_partenaire'] = Partenaire.generer_code(t)
        form = SuperAdminPartenaireForm(initial=initial)
    return render(request, 'superadmin/partenaire_form.html', {'form': form, 'obj': None})

@superadmin_required
def partenaire_detail(request, pk):
    partenaire = get_object_or_404(Partenaire.all_objects.filter(supprime=False), pk=pk)
    # Mutualités via FK direct + via RelationPartenaireMutualite
    from core.models import Mutualite
    mutualites_fk = Mutualite.objects.filter(partenaires=partenaire)
    mutualites_rel = Mutualite.objects.filter(relations_partenaires__partenaire=partenaire)
    mutualites = (mutualites_fk | mutualites_rel).distinct().order_by('-date_creation')
    
    return render(request, 'superadmin/partenaire_detail.html', {
        'partenaire': partenaire,
        'mutualites_liees': mutualites,
    })

@superadmin_required
def partenaire_modifier(request, pk):
    obj = get_object_or_404(Partenaire.all_objects.filter(supprime=False), pk=pk)
    if request.method == 'POST':
        form = SuperAdminPartenaireForm(request.POST, request.FILES, instance=obj)
        if form.is_valid():
            form.save()
            JournalAudit.objects.create(
                mutualite=obj.mutualite, utilisateur=request.user,
                action='MODIFICATION', module='superadmin',
                description=f"Modification partenaire : {obj.nom_etablissement}",
                objet_type='Partenaire', objet_id=obj.pk,
                adresse_ip=request.META.get('REMOTE_ADDR', ''),
            )
            messages.success(request, f"Partenaire {obj.nom_etablissement} modifié.")
            return redirect('superadmin:partenaire_detail', pk=obj.pk)
    else:
        form = SuperAdminPartenaireForm(instance=obj)
    return render(request, 'superadmin/partenaire_form.html', {'form': form, 'obj': obj})

@superadmin_required
def partenaire_suspendre(request, pk):
    partenaire = get_object_or_404(Partenaire.all_objects.filter(supprime=False), pk=pk)
    if partenaire.statut not in ('ACTIF', 'EN_ATTENTE'):
        messages.error(request, f"Impossible de suspendre. Statut actuel : {partenaire.get_statut_display()}")
        return redirect('superadmin:partenaire_detail', pk=pk)
    partenaire.statut = 'SUSPENDU'
    partenaire.save(update_fields=['statut'])
    JournalAudit.objects.create(
        mutualite=partenaire.mutualite, utilisateur=request.user,
        action='MODIFICATION', module='superadmin',
        description=f"Suspension partenaire : {partenaire.nom_etablissement}",
        objet_type='Partenaire', objet_id=partenaire.pk,
        adresse_ip=request.META.get('REMOTE_ADDR', ''),
    )
    messages.success(request, f"Partenaire {partenaire.nom_etablissement} suspendu.")
    return redirect('superadmin:partenaire_detail', pk=pk)

@superadmin_required
def partenaire_reactiver(request, pk):
    partenaire = get_object_or_404(Partenaire.all_objects.filter(supprime=False), pk=pk)
    if partenaire.statut != 'SUSPENDU':
        messages.error(request, "Seuls les partenaires suspendus peuvent être réactivés.")
        return redirect('superadmin:partenaire_detail', pk=pk)
    partenaire.statut = 'ACTIF'
    partenaire.save(update_fields=['statut'])
    JournalAudit.objects.create(
        mutualite=partenaire.mutualite, utilisateur=request.user,
        action='MODIFICATION', module='superadmin',
        description=f"Réactivation partenaire : {partenaire.nom_etablissement}",
        objet_type='Partenaire', objet_id=partenaire.pk,
        adresse_ip=request.META.get('REMOTE_ADDR', ''),
    )
    messages.success(request, f"Partenaire {partenaire.nom_etablissement} réactivé.")
    return redirect('superadmin:partenaire_detail', pk=pk)

@superadmin_required
def partenaire_bloquer(request, pk):
    partenaire = get_object_or_404(Partenaire.all_objects.filter(supprime=False), pk=pk)
    if partenaire.statut == 'BLOQUE':
        messages.error(request, "Ce partenaire est déjà bloqué.")
        return redirect('superadmin:partenaire_detail', pk=pk)
    partenaire.statut = 'BLOQUE'
    partenaire.save(update_fields=['statut'])
    JournalAudit.objects.create(
        mutualite=partenaire.mutualite, utilisateur=request.user,
        action='MODIFICATION', module='superadmin',
        description=f"Blocage partenaire : {partenaire.nom_etablissement}",
        objet_type='Partenaire', objet_id=partenaire.pk,
        adresse_ip=request.META.get('REMOTE_ADDR', ''),
    )
    messages.success(request, f"Partenaire {partenaire.nom_etablissement} bloqué.")
    return redirect('superadmin:partenaire_detail', pk=pk)

@superadmin_required
def partenaire_debloquer(request, pk):
    partenaire = get_object_or_404(Partenaire.all_objects.filter(supprime=False), pk=pk)
    if partenaire.statut != 'BLOQUE':
        messages.error(request, "Seuls les partenaires bloqués peuvent être débloqués.")
        return redirect('superadmin:partenaire_detail', pk=pk)
    partenaire.statut = 'ACTIF'
    partenaire.save(update_fields=['statut'])
    JournalAudit.objects.create(
        mutualite=partenaire.mutualite, utilisateur=request.user,
        action='MODIFICATION', module='superadmin',
        description=f"Déblocage partenaire : {partenaire.nom_etablissement}",
        objet_type='Partenaire', objet_id=partenaire.pk,
        adresse_ip=request.META.get('REMOTE_ADDR', ''),
    )
    messages.success(request, f"Partenaire {partenaire.nom_etablissement} débloqué.")
    return redirect('superadmin:partenaire_detail', pk=pk)

@superadmin_required
def partenaire_supprimer(request, pk):
    partenaire = get_object_or_404(Partenaire.all_objects.filter(supprime=False), pk=pk)
    if request.method == 'POST':
        nom = partenaire.nom_etablissement
        partenaire.soft_delete()
        JournalAudit.objects.create(
            mutualite=partenaire.mutualite, utilisateur=request.user,
            action='SUPPRESSION', module='superadmin',
            description=f"Suppression partenaire : {nom}",
            objet_type='Partenaire', objet_id=pk,
            adresse_ip=request.META.get('REMOTE_ADDR', ''),
        )
        messages.success(request, f"Partenaire {nom} supprimé.")
        return redirect('superadmin:partenaires_liste')
    return render(request, 'superadmin/partenaire_confirm_supprimer.html', {
        'partenaire': partenaire,
    })

@superadmin_required
def partenaire_utilisateurs(request, pk):
    partenaire = get_object_or_404(Partenaire.all_objects.filter(supprime=False), pk=pk)
    utilisateurs = partenaire.utilisateurs.all()

    if request.method == 'POST':
        form = SuperAdminUtilisateurPartenaireForm(request.POST)
        if form.is_valid():
            user = form.save(commit=False)
            user.partenaire = partenaire
            user.mutualite = partenaire.mutualite
            user.save()
            JournalAudit.objects.create(
                mutualite=partenaire.mutualite, utilisateur=request.user,
                action='CREATION', module='superadmin',
                description=f"Création utilisateur partenaire : {user.username} pour {partenaire.nom_etablissement}",
                objet_type='UtilisateurPartenaire', objet_id=user.pk,
                adresse_ip=request.META.get('REMOTE_ADDR', ''),
            )
            messages.success(request, f"Utilisateur {user.username} créé.")
            return redirect('superadmin:partenaire_utilisateurs', pk=partenaire.pk)
    else:
        form = SuperAdminUtilisateurPartenaireForm()

    return render(request, 'superadmin/partenaire_utilisateurs.html', {
        'partenaire': partenaire,
        'utilisateurs': utilisateurs,
        'form': form,
    })

@superadmin_required
def partenaire_utilisateur_modifier(request, partenaire_pk, user_pk):
    partenaire = get_object_or_404(Partenaire.all_objects.filter(supprime=False), pk=partenaire_pk)
    user = get_object_or_404(UtilisateurPartenaire, pk=user_pk, partenaire=partenaire)

    if request.method == 'POST':
        form = SuperAdminUtilisateurPartenaireForm(request.POST, instance=user)
        if form.is_valid():
            form.save()
            JournalAudit.objects.create(
                mutualite=partenaire.mutualite, utilisateur=request.user,
                action='MODIFICATION', module='superadmin',
                description=f"Modification utilisateur partenaire : {user.username}",
                objet_type='UtilisateurPartenaire', objet_id=user.pk,
                adresse_ip=request.META.get('REMOTE_ADDR', ''),
            )
            messages.success(request, f"Utilisateur {user.username} modifié.")
            return redirect('superadmin:partenaire_utilisateurs', pk=partenaire.pk)
    else:
        form = SuperAdminUtilisateurPartenaireForm(instance=user)

    return render(request, 'superadmin/partenaire_utilisateur_form.html', {
        'partenaire': partenaire,
        'user': user,
        'form': form,
    })

@superadmin_required
def partenaire_conventions(request, pk):
    partenaire = get_object_or_404(Partenaire.all_objects.filter(supprime=False), pk=pk)
    conventions = partenaire.conventions.select_related('mutualite').order_by('-date_debut')

    if request.method == 'POST':
        form = SuperAdminConventionForm(request.POST)
        if form.is_valid():
            convention = form.save(commit=False)
            convention.partenaire = partenaire
            convention.save()
            JournalAudit.objects.create(
                mutualite=convention.mutualite, utilisateur=request.user,
                action='CREATION', module='superadmin',
                description=f"Convention {convention.numero_convention} créée pour {partenaire.nom_etablissement}",
                objet_type='Convention', objet_id=convention.pk,
                adresse_ip=request.META.get('REMOTE_ADDR', ''),
            )
            messages.success(request, f"Convention {convention.numero_convention} créée.")
            return redirect('superadmin:partenaire_conventions', pk=partenaire.pk)
    else:
        form = SuperAdminConventionForm()

    return render(request, 'superadmin/partenaire_conventions.html', {
        'partenaire': partenaire,
        'conventions': conventions,
        'form': form,
    })

@superadmin_required
def partenaire_valider(request, pk):
    partenaire = get_object_or_404(Partenaire.all_objects.filter(supprime=False), pk=pk)
    if partenaire.statut != 'EN_ATTENTE':
        messages.warning(request, f"Le partenaire a déjà le statut {partenaire.get_statut_display()}.")
        return redirect('superadmin:partenaire_detail', pk=pk)

    partenaire.statut = 'ACTIF'
    partenaire.save(update_fields=['statut'])

    if not partenaire.conventions.filter(statut='ACTIVE').exists():
        from datetime import timedelta
        conv = Convention(
            partenaire=partenaire,
            mutualite=partenaire.mutualite,
            numero_convention=Convention._generer_numero(Convention),
            date_debut=timezone.now().date(),
            date_fin=timezone.now().date() + timedelta(days=365),
            statut='ACTIVE',
            created_by=request.user,
        )
        conv.save()

    JournalAudit.objects.create(
        mutualite=partenaire.mutualite, utilisateur=request.user,
        action='VALIDATION', module='superadmin',
        description=f"Validation partenaire : {partenaire.nom_etablissement}",
        objet_type='Partenaire', objet_id=partenaire.pk,
        adresse_ip=request.META.get('REMOTE_ADDR', ''),
    )
    messages.success(request, f"Partenaire {partenaire.nom_etablissement} validé.")
    return redirect('superadmin:partenaire_detail', pk=pk)

@superadmin_required
def partenaire_supprimer_validation(request, pk, mutualite_pk):
    partenaire = get_object_or_404(Partenaire.all_objects.filter(supprime=False), pk=pk)
    relation = RelationPartenaireMutualite.objects.filter(
        partenaire=partenaire, mutualite_id=mutualite_pk
    ).first()
    if relation:
        relation.delete()
    partenaire.statut = 'EN_ATTENTE'
    partenaire.mutualite = None
    partenaire.save(update_fields=['statut', 'mutualite'])
    JournalAudit.objects.create(
        utilisateur=request.user, action='MODIFICATION', module='superadmin',
        description=f"Validation supprimée pour {partenaire.nom_etablissement} (mutualité #{mutualite_pk})",
        objet_type='Partenaire', objet_id=partenaire.pk,
        adresse_ip=request.META.get('REMOTE_ADDR', ''),
    )
    messages.success(request, f"Validation retirée pour {partenaire.nom_etablissement}.")
    return redirect('superadmin:partenaire_detail', pk=pk)


@superadmin_required
def notifications_liste(request):
    from core.models import Notification
    mutualite_id = request.GET.get('mutualite')
    qs = Notification.objects.all()
    if mutualite_id:
        qs = qs.filter(mutualite_id=mutualite_id)
    qs = qs.select_related('mutualite', 'affilie')

    paginator = Paginator(qs, 20)
    page = paginator.get_page(request.GET.get('page'))

    mutualites = Mutualite.objects.all()
    context = {
        'notifications': page,
        'paginator': paginator,
        'mutualites': mutualites,
        'filtre_mutualite': int(mutualite_id) if mutualite_id else None,
    }
    return render(request, 'superadmin/notification_liste.html', context)


@superadmin_required
def validations_mutualites(request):
    q = request.GET.get('q', '')
    statut_filter = request.GET.get('statut', '')

    mutualites = Mutualite.objects.all().select_related('created_by').order_by('-date_creation')

    if q:
        mutualites = mutualites.filter(
            Q(nom__icontains=q) | Q(code_mutualite__icontains=q) |
            Q(nif__icontains=q) | Q(nom_responsable__icontains=q)
        )

    if statut_filter:
        mutualites = mutualites.filter(statut=statut_filter)

    paginator = Paginator(mutualites, 5)
    page_number = request.GET.get('page')
    mutualites_page = paginator.get_page(page_number)

    # Stats for badges
    stats = Mutualite.objects.values('statut').annotate(count=Count('pk'))

    return render(request, 'superadmin/validations_mutualites.html', {
        'mutualites': mutualites_page,
        'q': q,
        'statut_filter': statut_filter,
        'paginator': paginator,
        'stats': stats,
    })


@superadmin_required
def valider_mutualite(request, pk):
    mutualite = get_object_or_404(Mutualite, pk=pk, statut='EN_ATTENTE')
    admin_user = User.objects.filter(mutualite=mutualite, role='ADMIN_MUTUALITE').first()
    
    # Partenaires liés avec pagination (5 par page)
    from django.core.paginator import Paginator
    partenaires_qs = mutualite.partenaires.all().order_by('-date_creation')
    page_partenaire = request.GET.get('page_partenaire', 1)
    paginator_partenaires = Paginator(partenaires_qs, 5)
    partenaires_page = paginator_partenaires.get_page(page_partenaire)

    if request.method == 'POST':
        action = request.POST.get('action', 'valider')
        
        if action == 'supprimer':
            nom = mutualite.nom
            # Supprimer l'admin associé aussi
            if admin_user:
                admin_user.delete()
            mutualite.delete()
            JournalAudit.objects.create(
                utilisateur=request.user, action='SUPPRESSION', module='superadmin',
                description=f"Mutualité en attente '{nom}' supprimée par {request.user.username}",
                objet_type='Mutualite', objet_id=pk,
                adresse_ip=request.META.get('REMOTE_ADDR', ''),
            )
            messages.success(request, f"Mutualité « {nom} » supprimée définitivement.")
            return redirect('superadmin:validations_mutualites')

        type_plan = request.POST.get('type_plan')
        if not type_plan:
            messages.error(request, 'Veuillez sélectionner un type de plan.')
            return redirect('superadmin:valider_mutualite', pk=pk)

        with transaction.atomic():
            mutualite.activer_essai(cree_par=request.user)

            if admin_user and not admin_user.actif:
                admin_user.actif = True
                admin_user.save(update_fields=['actif'])

            JournalAudit.objects.create(
                mutualite=mutualite, utilisateur=request.user,
                action='VALIDATION', module='superadmin',
                description=f"Mutualité '{mutualite.nom}' validée et essai activé par {request.user.username}",
                objet_type='Mutualite', objet_id=mutualite.pk,
                adresse_ip=request.META.get('REMOTE_ADDR', ''),
            )

        messages.success(
            request,
            f"Mutualité « {mutualite.nom} » activée (essai 14 jours). "
            f"Contactez l'administrateur au {mutualite.telephone} pour l'informer."
        )
        return redirect('superadmin:validations_mutualites')

    return render(request, 'superadmin/valider_mutualite.html', {
        'mutualite': mutualite,
        'admin_user': admin_user,
        'partenaires': partenaires_page,
        'paginator_partenaires': paginator_partenaires,
    })


@superadmin_required
def validations_partenaires(request):
    q = request.GET.get('q', '')
    statut_filter = request.GET.get('statut', '')

    partenaires = Partenaire.objects.all().select_related('mutualite', 'created_by').order_by('-date_creation')

    if q:
        partenaires = partenaires.filter(
            Q(nom_etablissement__icontains=q) | Q(code_partenaire__icontains=q) |
            Q(nif__icontains=q) | Q(nom_responsable__icontains=q)
        )

    if statut_filter:
        partenaires = partenaires.filter(statut=statut_filter)

    paginator = Paginator(partenaires, 5)
    page_number = request.GET.get('page')
    partenaires_page = paginator.get_page(page_number)

    # Stats for badges
    stats = Partenaire.objects.values('statut').annotate(count=Count('pk'))

    return render(request, 'superadmin/validations_partenaires.html', {
        'partenaires': partenaires_page,
        'q': q,
        'statut_filter': statut_filter,
        'paginator': paginator,
        'stats': stats,
    })


@superadmin_required
def valider_partenaire(request, pk):
    partenaire = get_object_or_404(Partenaire, pk=pk, statut='EN_ATTENTE')
    mutualites = Mutualite.objects.filter(
        statut__in=['ACTIF', 'ESSAI']
    ).order_by('nom')

    if request.method == 'POST':
        mutualite_id = request.POST.get('mutualite')
        if not mutualite_id:
            messages.error(request, 'Veuillez sélectionner une mutualité pour ce partenaire.')
            return redirect('superadmin:valider_partenaire', pk=pk)

        mutualite = get_object_or_404(Mutualite, pk=mutualite_id)

        with transaction.atomic():
            partenaire.statut = 'ACTIF'
            partenaire.save(update_fields=['statut'])

            RelationPartenaireMutualite.objects.create(
                partenaire=partenaire,
                mutualite=mutualite,
                statut='ACTIVE',
                date_demande=timezone.now(),
            )

            conv = Convention(
                partenaire=partenaire,
                mutualite=mutualite,
                numero_convention=Convention._generer_numero(Convention),
                date_debut=timezone.now().date(),
                date_fin=timezone.now().date() + timedelta(days=365),
                statut='ACTIVE',
                created_by=request.user,
            )
            conv.save()

            JournalAudit.objects.create(
                mutualite=mutualite, utilisateur=request.user,
                action='VALIDATION', module='superadmin',
                description=f"Validation partenaire '{partenaire.nom_etablissement}' lié à {mutualite.nom}",
                objet_type='Partenaire', objet_id=partenaire.pk,
                adresse_ip=request.META.get('REMOTE_ADDR', ''),
            )

        messages.success(
            request,
            f"Partenaire « {partenaire.nom_etablissement} » activé et lié à « {mutualite.nom} ». "
            f"Contactez le responsable au {partenaire.telephone} pour l'informer."
        )
        return redirect('superadmin:validations_partenaires')

    return render(request, 'superadmin/valider_partenaire.html', {
        'partenaire': partenaire,
        'mutualites': mutualites,
    })
